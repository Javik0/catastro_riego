# -*- coding: utf-8 -*-
"""
Listado de titulares entrevistados y de quienes dicen conocer el proyecto.

Lo pide el sociólogo (Alexis Guerrero, 1-oct-2026): la lista de todos los
catastrados, con cédula y comunidad, y de ellos los que contestaron que sí
conocen el proyecto, para poder sostener con documentos que los titulares
existen, fueron entrevistados y contestaron.

Qué se cuenta y por qué
-----------------------
Se parte de las FICHAS PRINCIPALES (4.307): cada una es la entrevista a un
titular. Las adicionales son otros terrenos del mismo titular y no son
personas nuevas.

La cifra del 87 % que circula sale del informe del sociólogo, donde la
columna «Conoce la presa %» se calcula sobre quienes RESPONDIERON la pregunta
(Sí + No = 4.132), no sobre las 4.307. Sobre las 4.307 es 83,7 %; y quienes
contestaron son el 95,9 %. El Excel trae las tres lecturas para que no se
mezclen.

«Existen» se apoya en la cédula, y solo hasta donde el dato alcanza: el
dígito verificador comprueba que el número está bien formado, NO que figure
en el Registro Civil. Y 4.307 fichas no son 4.307 personas: 217 cédulas
aparecen en más de una ficha principal. La hoja PARA REVISAR lista esos
casos, los de cédula inválida y los que no la tienen, para que campo los
resuelva antes de que alguien más los encuentre.

El campo `consentimiento_inform` está vacío en las 4.307: no hay
consentimiento informado registrado en el dato y por eso no se usa como
evidencia.

Uso:  python -X utf8 scripts/generar_listado_titulares.py
"""
import collections
import io
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

from armar_entrega_final import catalogo, CONSTANTS  # noqa: E402
from corregir_cedulas_homonimos import cedula_valida  # noqa: E402
from informe_estilo import FECHA_CORTE  # noqa: E402

BASE = os.path.abspath(os.path.join(AQUI, '..'))
GEO = os.path.join(BASE, 'public', 'geo')
SALIDA = (r'C:\Users\HP\OneDrive\Escritorio'
          r'\LISTADO DE TITULARES ENTREVISTADOS - conocimiento del proyecto.xlsx')

AZUL = '1E3A8A'
FILL_CAB = PatternFill('solid', fgColor=AZUL)
FILL_TOT = PatternFill('solid', fgColor='EEF3EF')
LADO = Side(style='thin', color='CBD5E1')
BORDE = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)
CAB = Font(bold=True, color='FFFFFF', size=10)


def nombre_tecnicos():
    """usuario de campo -> nombre, leído de constants.ts: la misma fuente
    que usa la web, para que no haya dos tablas que se separen."""
    s = io.open(CONSTANTS, encoding='utf-8').read()
    return dict(re.findall(r"'([^']+)':\s*\{\s*nombre:\s*'([^']+)'", s))


def fecha(iso):
    try:
        return datetime.strptime(str(iso)[:10], '%Y-%m-%d').strftime('%d/%m/%Y')
    except Exception:
        return ''


def estado_cedula(c):
    if not c:
        return 'Sin cédula'
    r = cedula_valida(c)
    if r is True:
        return 'Dígito verificador válido'
    if r is False:
        return 'Dígito verificador NO válido'
    return 'Otro formato (RUC / pasaporte)'


def cargar():
    with io.open(os.path.join(GEO, 'fichas_predios.geojson'), encoding='utf-8') as f:
        fichas = [ft['properties'] for ft in json.load(f)['features']]
    with io.open(os.path.join(GEO, 'codificacion_fichas.json'), encoding='utf-8') as f:
        cod = json.load(f)['fichas']
    _, por_numero = catalogo()
    tecnicos = nombre_tecnicos()

    pri = [p for p in fichas if not p.get('es_ficha_hija')]
    por_cedula = collections.defaultdict(list)
    for p in pri:
        c = (p.get('cedula') or '').strip()
        if c:
            por_cedula[c].append(cod[p['id']])

    filas = []
    for p in pri:
        codigo = cod[p['id']]
        ced = (p.get('cedula') or '').strip()
        # El sector y la comunidad salen del CÓDIGO de la ficha, no del texto
        # de la comunidad: el código no depende de cómo se escribió el nombre.
        sec, com = int(codigo[1:3]), int(codigo[5:7])
        resp = str(p.get('conoce_presa') or '').strip()
        nom = f"{p.get('apellidos') or ''} {p.get('nombres') or ''}".strip() \
            or (p.get('propietario') or '').strip()
        otras = [c for c in por_cedula.get(ced, []) if c != codigo] if ced else []
        filas.append({
            'codigo': codigo, 'sector': f'Sector {sec}', 'com_n': com,
            'comunidad': por_numero.get(com, f'Comunidad {com}'),
            'cedula': ced, 'estado_ced': estado_cedula(ced), 'nombre': nom,
            'clave': (p.get('clave_catastral') or '').strip(),
            'conoce': resp if resp in ('Sí', 'No') else 'Sin respuesta',
            'fecha': fecha(p.get('fecha_creacion')),
            'tecnico': tecnicos.get(p.get('creado_por'), p.get('creado_por') or ''),
            'gps': bool(p.get('coord_x_utm') and p.get('coord_y_utm')),
            'otras': otras,
        })
    filas.sort(key=lambda r: r['codigo'])
    return filas, por_cedula


def escribir_hoja(ws, cabeceras, anchos, filas, texto=()):
    ws.append(cabeceras)
    for i in range(1, len(cabeceras) + 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for fila in filas:
        ws.append(fila)
    for r in range(2, ws.max_row + 1):
        for i in range(1, len(cabeceras) + 1):
            c = ws.cell(row=r, column=i)
            c.border = BORDE
            c.font = Font(size=9.5)
            if i in texto:
                # cédula y clave como TEXTO: Excel las trataría como número y
                # las pasaría a notación científica.
                c.number_format = '@'
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cabeceras))}{ws.max_row}'
    ws.row_dimensions[1].height = 30


def casos_para_revisar(filas):
    """Las fichas cuya cédula hay que confirmar, con el motivo en texto.

    Una sola definición para todo lo que la use —la hoja PARA REVISAR de este
    Excel y la solicitud a campo (`generar_verificacion_cedulas.py`)—: si dos
    sitios decidieran por separado qué fichas son «a revisar», las cuentas se
    separarían sin que nadie lo notara.
    """
    revisar = []
    for x in filas:
        motivos = []
        if x['estado_ced'] == 'Sin cédula':
            motivos.append('Sin cédula registrada')
        elif x['estado_ced'] == 'Dígito verificador NO válido':
            motivos.append('Dígito verificador no válido')
        if x['otras']:
            motivos.append(f"Misma cédula en {len(x['otras']) + 1} fichas principales")
        if motivos:
            revisar.append((x, '; '.join(motivos)))
    return revisar


def main():
    filas, por_cedula = cargar()
    N = len(filas)
    si = [r for r in filas if r['conoce'] == 'Sí']
    no = [r for r in filas if r['conoce'] == 'No']
    sr = [r for r in filas if r['conoce'] == 'Sin respuesta']
    resp = len(si) + len(no)

    def pct(a, b):
        return (a / b) if b else 0

    wb = Workbook()

    # ── RESUMEN ───────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'RESUMEN'
    ws.sheet_view.showGridLines = False
    ws['A1'] = 'Titulares entrevistados y conocimiento del proyecto'
    ws['A1'].font = Font(bold=True, size=15, color=AZUL)
    ws['A2'] = ('Padrón de usuarios · Sistema de riego comunitario '
                f'Guanguilquí–Porotog · Información al {FECHA_CORTE}')
    ws['A2'].font = Font(size=10, color='64748B')

    def bloque(fila, titulo, cabeceras, datos, formatos):
        ws.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=11, color=AZUL)
        for j, h in enumerate(cabeceras, 1):
            c = ws.cell(row=fila + 1, column=j, value=h)
            c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
            c.alignment = Alignment(horizontal='center', wrap_text=True)
        r = fila + 2
        for d in datos:
            for j, v in enumerate(d, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = BORDE
                if j > 1 and v is not None:
                    c.number_format = formatos[j - 2]
                    c.alignment = Alignment(horizontal='right')
            r += 1
        return r + 1

    p0 = '0.0%'
    n0 = '#,##0'
    r = 4
    r = bloque(r, '1. ¿Cuántos contestaron?',
               ['Concepto', 'Fichas', '% de las fichas', '% de quienes respondieron'],
               [['Fichas principales (una entrevista por titular)', N, 1.0, None],
                ['Respondieron «¿Conoce el proyecto?»', resp, pct(resp, N), 1.0],
                ['   Dicen que SÍ conocen el proyecto', len(si), pct(len(si), N), pct(len(si), resp)],
                ['   Dicen que NO', len(no), pct(len(no), N), pct(len(no), resp)],
                ['Sin respuesta registrada', len(sr), pct(len(sr), N), None]],
               [n0, p0, p0])
    con_fecha = sum(1 for x in filas if x['fecha'])
    con_tec = sum(1 for x in filas if x['tecnico'])
    con_gps = sum(1 for x in filas if x['gps'])
    r = bloque(r, '2. ¿Fueron entrevistados? Lo que consta en cada ficha',
               ['Evidencia', 'Fichas', '% de las fichas'],
               [['Con fecha de levantamiento', con_fecha, pct(con_fecha, N)],
                ['Con técnico investigador registrado', con_tec, pct(con_tec, N)],
                ['Con coordenadas del punto de levantamiento', con_gps, pct(con_gps, N)]],
               [n0, p0])
    ced = collections.Counter(x['estado_ced'] for x in filas)
    con_ced = N - ced['Sin cédula']
    r = bloque(r, '3. ¿Existen? La cédula de cada titular',
               ['Cédula', 'Fichas', '% de las fichas'],
               [['Con cédula registrada', con_ced, pct(con_ced, N)],
                ['   con dígito verificador válido', ced['Dígito verificador válido'],
                 pct(ced['Dígito verificador válido'], N)],
                ['   con dígito verificador NO válido', ced['Dígito verificador NO válido'],
                 pct(ced['Dígito verificador NO válido'], N)],
                ['   en otro formato (RUC / pasaporte)', ced['Otro formato (RUC / pasaporte)'],
                 pct(ced['Otro formato (RUC / pasaporte)'], N)],
                ['Sin cédula registrada', ced['Sin cédula'], pct(ced['Sin cédula'], N)]],
               [n0, p0])
    dup = {c: v for c, v in por_cedula.items() if len(v) > 1}
    r = bloque(r, '4. ¿Cuántas personas distintas son?',
               ['Concepto', 'Cantidad'],
               [['Cédulas distintas entre las fichas principales', len(por_cedula)],
                ['Fichas principales sin cédula (no se pueden contar como distintas)',
                 ced['Sin cédula']],
                ['Cédulas que aparecen en más de una ficha principal', len(dup)],
                ['   fichas principales que comparten cédula',
                 sum(len(v) for v in dup.values())]],
               [n0])
    notas = [
        'El dígito verificador comprueba que el número de cédula está bien formado; no certifica '
        'que figure en el Registro Civil.',
        'Las fichas principales son entrevistas, no personas: una misma cédula puede tener varias '
        'fichas principales (la hoja PARA REVISAR las lista).',
        'El campo de consentimiento informado está vacío en todas las fichas; no se usa como evidencia.',
        'El 87 % del informe del sociólogo se calcula sobre quienes respondieron '
        '(fila «Dicen que SÍ», última columna).',
    ]
    ws.cell(row=r, column=1, value='Cómo se lee').font = Font(bold=True, size=11, color=AZUL)
    for k, t in enumerate(notas, 1):
        c = ws.cell(row=r + k, column=1, value=f'• {t}')
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r + k, start_column=1, end_row=r + k, end_column=4)
        ws.row_dimensions[r + k].height = 30
    for col, w in zip('ABCD', (62, 14, 16, 24)):
        ws.column_dimensions[col].width = w

    # ── LISTADO COMPLETO y CONOCEN EL PROYECTO ───────────────────────────
    cab = ['N°', 'Código de ficha', 'Sector', 'Comunidad', 'Cédula', 'Apellidos y nombres',
           '¿Conoce el proyecto?', 'Clave catastral', 'Fecha de levantamiento', 'Técnico']
    anchos = [7, 18, 10, 30, 14, 44, 16, 17, 14, 22]

    def fila_lista(i, x):
        return [i, x['codigo'], x['sector'], x['comunidad'], x['cedula'], x['nombre'],
                x['conoce'], x['clave'], x['fecha'], x['tecnico']]

    escribir_hoja(wb.create_sheet('LISTADO COMPLETO'), cab, anchos,
                  [fila_lista(i, x) for i, x in enumerate(filas, 1)], texto=(5, 8))
    escribir_hoja(wb.create_sheet('CONOCEN EL PROYECTO'), cab, anchos,
                  [fila_lista(i, x) for i, x in enumerate(si, 1)], texto=(5, 8))

    # ── POR COMUNIDAD ────────────────────────────────────────────────────
    por_com = collections.OrderedDict()
    for x in sorted(filas, key=lambda x: (x['com_n'], x['codigo'])):
        d = por_com.setdefault((x['com_n'], x['comunidad'], x['sector']),
                               {'n': 0, 'si': 0, 'no': 0, 'sr': 0})
        d['n'] += 1
        d['si' if x['conoce'] == 'Sí' else 'no' if x['conoce'] == 'No' else 'sr'] += 1
    ws = wb.create_sheet('POR COMUNIDAD')
    cab_c = ['N°', 'Comunidad', 'Sector', 'Fichas principales', 'Respondieron', 'Sí conocen',
             'No conocen', 'Sin respuesta', '% que respondió',
             '% Sí (sobre quienes respondieron)', '% Sí (sobre las fichas)']
    filas_c = []
    for (n, nom, sec), d in por_com.items():
        rs = d['si'] + d['no']
        filas_c.append([n, nom, sec, d['n'], rs, d['si'], d['no'], d['sr'],
                        pct(rs, d['n']), pct(d['si'], rs), pct(d['si'], d['n'])])
    escribir_hoja(ws, cab_c, [6, 32, 10, 12, 12, 11, 11, 12, 13, 18, 16], filas_c)
    for rr in range(2, ws.max_row + 1):
        for cc in (9, 10, 11):
            ws.cell(row=rr, column=cc).number_format = p0
    ws.append(['', 'TOTAL', '', N, resp, len(si), len(no), len(sr),
               pct(resp, N), pct(len(si), resp), pct(len(si), N)])
    for cc in range(1, 12):
        c = ws.cell(row=ws.max_row, column=cc)
        c.font, c.fill, c.border = Font(bold=True, size=10), FILL_TOT, BORDE
        if cc in (9, 10, 11):
            c.number_format = p0

    # ── PARA REVISAR ─────────────────────────────────────────────────────
    revisar = casos_para_revisar(filas)
    cab_r = ['N°', 'Código de ficha', 'Sector', 'Comunidad', 'Cédula', 'Apellidos y nombres',
             '¿Conoce el proyecto?', 'Qué revisar', 'Otras fichas con la misma cédula']
    filas_r = [[i, x['codigo'], x['sector'], x['comunidad'], x['cedula'], x['nombre'],
                x['conoce'], m,
                ', '.join(x['otras'][:6]) + (' …' if len(x['otras']) > 6 else '')]
               for i, (x, m) in enumerate(revisar, 1)]
    escribir_hoja(wb.create_sheet('PARA REVISAR'), cab_r,
                  [7, 18, 10, 30, 14, 44, 16, 44, 60], filas_r, texto=(5,))

    wb.save(SALIDA)

    # Se vuelve a leer el archivo y se cuadran las cuentas: lo que se entrega
    # es lo que quedó en disco, no lo que el programa cree haber escrito.
    from openpyxl import load_workbook
    chk = load_workbook(SALIDA, read_only=True)
    n_lista = chk['LISTADO COMPLETO'].max_row - 1
    n_si = chk['CONOCEN EL PROYECTO'].max_row - 1
    suma_n = suma_si = 0
    for row in chk['POR COMUNIDAD'].iter_rows(min_row=2, values_only=True):
        if row[1] != 'TOTAL':
            suma_n += row[3]
            suma_si += row[5]
    assert n_lista == N, (n_lista, N)
    assert n_si == len(si), (n_si, len(si))
    assert suma_n == N and suma_si == len(si), (suma_n, suma_si)

    def pt(v):
        return f'{v:,}'.replace(',', '.')

    print(f'Fichas principales : {pt(N)}')
    print(f'  respondieron     : {pt(resp)} ({pct(resp, N) * 100:.1f} %)')
    print(f'  dicen que SÍ     : {pt(len(si))} ({pct(len(si), N) * 100:.1f} % de las fichas · '
          f'{pct(len(si), resp) * 100:.1f} % de quienes respondieron)')
    print(f'  hoja PARA REVISAR: {len(revisar)} fichas')
    print('Cuentas cuadradas tras releer el archivo: listado, Sí y suma por comunidad.')
    print('✔', SALIDA)
    return 0


if __name__ == '__main__':
    sys.exit(main())
