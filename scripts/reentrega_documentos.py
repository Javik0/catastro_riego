# -*- coding: utf-8 -*-
"""
Documentos de la REENTREGA v2 del Producto 5 (fase C del plan de subsanación).
Los llama `armar_reentrega_producto5.py`; no se ejecuta solo.

  diccionario_v2   Diccionario de datos leído del GeoPackage v2: campos, tipo,
                   nombre en el Shapefile, dominio y vacíos; universos (C03) y
                   superficies (9-a). PDF y Excel.
  nota_avance      Nota aclaratoria al documento 6 «Porcentaje de avance y
                   fecha de corte» (C05). El documento original no se toca.
  guia_v2          «Qué cambia en la versión 2»: qué se agregó, qué observación
                   atiende y la constancia de que ninguna cifra cambió.
  inventario       Cada archivo de la reentrega, si es nuevo, actualizado o igual
                   a la v1, y qué observación atiende (7.1-a, 11.1-b).
"""
import collections
import hashlib
import json
import os
import sqlite3
import struct
import zipfile

import armar_producto5 as ap5
import generar_diccionario_datos as dic1
from generar_fichas_pdf import Paragraph, Spacer, mm, tabla_datos, titulo_seccion
from informe_estilo import FECHA_CORTE, esn

GEO = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'public', 'geo')

# ── campos y tablas nuevas de la v2 ────────────────────────────────────────
CAPAS_NUEVAS = [
    ('correspondencia_fichas', 'Tabla',
     'Relación ficha – titular – predio – polígono, una fila por ficha. Ningún registro se elimina: lo que está en '
     'revisión se marca.'),
    ('fichas_sin_poligono', 'Tabla', 'Fichas principales cuyo predio no tiene polígono en el catastro del cantón.'),
    ('catalogo_comunidades', 'Tabla', 'Las 50 comunidades: código único, nombre oficial y nombre usado en la base.'),
    ('registro_cambios', 'Tabla',
     'Altas, bajas y cambios ficha por ficha desde el 18 de agosto de 2026 (antes de la depuración) hasta la base '
     'entregada.'),
    ('metadatos', 'Tabla', 'Fecha de corte, versión y cambios respecto de la entrega anterior.'),
]
DESC_V2 = {
    'predios_investigados': {
        'sistema_riego_declarado': 'Sistema u organización de riego que declararon los titulares del predio. No es el '
                                   'sector 1–3 (ese está en «sector»). En la v1 se llamaba sector_riego.',
        'codigos_ficha': 'Códigos de las fichas levantadas sobre el predio (S01-C22-R001-F01). Enlazan el predio '
                         'con la capa fichas y con los PDF de las fichas.',
        'codigo_comunidad': 'Código de la comunidad en el catálogo (C01 a C50).',
        'comunidad_oficial': 'Nombre oficial de la comunidad según el catálogo.',
        'sector': 'Sector de investigación (Sector 1, 2 o 3) según el catálogo de comunidades.',
        'fecha_corte': 'Fecha de corte de la información.',
        'version': 'Versión de la entrega.',
        'cultivos_predio': 'Cultivos declarados en las fichas del predio, con su superficie en hectáreas, en texto. '
                           'Es un resumen para leer; la fuente de cálculo es la tabla cultivos.',
        'animales_predio': 'Especies pecuarias declaradas en las fichas del predio, con sus cabezas, en texto. Es un '
                           'resumen para leer; la fuente de cálculo es la tabla animales.',
    },
    'fichas': {
        'sector': 'Sistema u organización de riego declarado por el titular. No es el sector 1–3: ese está en '
                  'sector_catalogo.',
        'codigo_comunidad': 'Código de la comunidad en el catálogo (C01 a C50).',
        'comunidad_oficial': 'Nombre oficial de la comunidad según el catálogo.',
        'sector_catalogo': 'Sector de investigación (Sector 1, 2 o 3) según el catálogo de comunidades.',
    },
    'cultivos': {
        'cultivo': 'Cultivo declarado. Si el técnico registró además un detalle en el campo «otro», manda el tipo de '
                   'cultivo y el detalle va en cultivo_detalle.',
        'cultivo_detalle': 'Detalle que el técnico escribió en el campo «otro» cuando ya había elegido un tipo de '
                           'cultivo (por ejemplo, «Aguacate» en Frutales).',
    },
    'correspondencia_fichas': {
        'codigo_ficha': 'Código único de la ficha (S01-C22-R001-F01).',
        'ficha_id': 'Identificador interno de la ficha (UUID de QField).',
        'tipo_ficha': 'Ficha principal o ficha adicional.',
        'codigo_ficha_principal': 'En una ficha adicional, el código de la ficha principal que la declaró.',
        'titular': 'Apellidos y nombres del titular de la ficha.',
        'cedula': 'Cédula (o RUC) registrada en la ficha.',
        'codigo_comunidad': 'Código de la comunidad (C01 a C50).',
        'comunidad_oficial': 'Nombre oficial de la comunidad.',
        'sector': 'Sector de investigación (Sector 1, 2 o 3).',
        'clave_catastral': 'Clave catastral del predio de la ficha.',
        'tiene_poligono': 'Sí, si el predio está en el catastro del cantón Cayambe; No, en caso contrario.',
        'marcas_revision': 'Revisiones en curso: ficha sin polígono y su causa, posible registro duplicado de la misma '
                           'entrevista o cédula en verificación de campo. Vacío = sin observación.',
    },
    'fichas_sin_poligono': {
        'codigo_ficha': 'Código único de la ficha.', 'titular': 'Titular de la ficha.',
        'comunidad_oficial': 'Comunidad de la ficha.', 'clave_declarada': 'Clave catastral que declaró el titular.',
        'area_declarada_m2': 'Superficie declarada en la ficha, en metros cuadrados.',
        'causa': 'Por qué el predio no tiene polígono en el catastro del cantón.',
    },
    'catalogo_comunidades': {
        'codigo_comunidad': 'Código único de la comunidad (C01 a C50).', 'numero': 'Número oficial de la comunidad.',
        'comunidad_oficial': 'Nombre oficial.', 'nombre_en_la_base': 'Nombre con que la comunidad figura en las capas.',
        'sector': 'Sector de investigación.', 'fichas': 'Fichas de la comunidad.',
        'predios_con_poligono': 'Predios de la comunidad con polígono en el catastro.',
    },
    'registro_cambios': {
        'codigo_ficha': 'Código de la ficha.', 'ficha_id': 'Identificador interno de la ficha.',
        'titular': 'Titular de la ficha.', 'tipo_cambio': 'Alta, baja o modificación.',
        'campo': 'Dato que cambió (cultivos y animales se resumen en registros y superficie o cabezas).',
        'valor_antes': 'Valor al 18 de agosto de 2026, antes de la depuración.',
        'valor_despues': 'Valor en la base entregada.',
    },
    'metadatos': {'elemento': 'Qué se describe.', 'valor': 'Valor.'},
}
CATEGORICOS_V2 = set(dic1.CATEGORICOS) | {
    ('predios_investigados', 'sistema_riego_declarado'), ('predios_investigados', 'sector'),
    ('fichas', 'sector_catalogo'), ('correspondencia_fichas', 'tipo_ficha'),
    ('correspondencia_fichas', 'tiene_poligono'), ('correspondencia_fichas', 'sector'),
    ('registro_cambios', 'tipo_cambio'), ('registro_cambios', 'campo'), ('catalogo_comunidades', 'sector'),
}


def _campos_dbf(ruta):
    if not os.path.exists(ruta):
        return []
    b = open(ruta, 'rb').read(1 << 16)
    hl = struct.unpack('<H', b[8:10])[0]
    return [b[i:i + 11].split(b'\0')[0].decode('latin-1') for i in range(32, hl - 1, 32)]


def estructura_v2(gpkg, shp_dir):
    con = sqlite3.connect(gpkg)
    out = []
    for capa, geom, proposito in list(dic1.CAPAS) + CAPAS_NUEVAS:
        cols = [r for r in con.execute(f'pragma table_info("{capa}")') if r[1] not in ('fid', 'geom')]
        en_shp = dict(zip([r[1] for r in cols], _campos_dbf(os.path.join(shp_dir, capa + '.dbf'))))
        total = con.execute(f'select count(*) from "{capa}"').fetchone()[0]
        campos = []
        for r in cols:
            nombre, tipo = r[1], r[2]
            dominio = ''
            if (capa, nombre) in CATEGORICOS_V2:
                vals = [x[0] for x in con.execute(f'select distinct "{nombre}" from "{capa}" where "{nombre}" is not '
                                                  f'null and trim("{nombre}")<>\'\' order by 1')]
                dominio = ' · '.join(str(v) for v in vals) if len(vals) <= 12 else f'{esn(len(vals))} valores distintos'
            elif nombre in ('agua_consumo', 'energia_electrica'):
                dominio = dic1.SI_NO_VACIO
            vacios = con.execute(f'select count(*) from "{capa}" where "{nombre}" is null or trim("{nombre}")=\'\'') \
                .fetchone()[0]
            desc = DESC_V2.get(capa, {}).get(nombre) or dic1.DESCRIPCIONES.get(capa, {}).get(nombre, '')
            campos.append(dict(campo=nombre, tipo=dic1.TIPO_LEGIBLE.get(tipo, tipo or 'Texto'), descripcion=desc,
                               dominio=dominio, shp=en_shp.get(nombre, '—'),
                               vacios=f'{esn(vacios)} de {esn(total)}' if vacios else '0'))
        out.append(dict(capa=capa, geometria=geom, proposito=proposito, registros=total, campos=campos))
    con.close()
    return out


def universos(gpkg):
    con = sqlite3.connect(gpkg)
    tipo = dict(con.execute('select tipo_ficha, count(*) from fichas group by 1').fetchall())
    ced = con.execute("select count(distinct trim(cedula)) from fichas where tipo_ficha='Ficha principal' and "
                      "trim(coalesce(cedula,''))<>''").fetchone()[0]
    sin_ced = con.execute("select count(*) from fichas where tipo_ficha='Ficha principal' and "
                          "trim(coalesce(cedula,''))=''").fetchone()[0]
    pred, cat_m2, decl_m2 = con.execute('select count(*), sum(area_catastro_m2), sum(area_declarada_m2) '
                                        'from predios_investigados').fetchone()
    tot_m2, rie_m2, sin_m2 = con.execute('select sum(area_total_m2), sum(area_riego_m2), sum(area_sin_riego_m2) '
                                         'from fichas').fetchone()
    sinpol = con.execute('select count(*) from fichas_sin_poligono').fetchone()[0]
    coms = con.execute('select count(*) from catalogo_comunidades where fichas>0').fetchone()[0]
    con.close()
    tot = json.load(open(os.path.join(GEO, 'superficie_por_comunidad.json'), encoding='utf-8'))['total']
    return dict(principales=tipo.get('Ficha principal', 0), adicionales=tipo.get('Ficha adicional', 0),
                cedulas=ced, sin_cedula=sin_ced, predios=pred, cat_ha=cat_m2 / 1e4, decl_pol_ha=decl_m2 / 1e4,
                decl_ha=tot_m2 / 1e4, riego_ha=rie_m2 / 1e4, sin_riego_ha=sin_m2 / 1e4, sinpol=sinpol,
                comunidades=coms, riego_ajustado_ha=tot['riego_ajustado_ha'],
                sin_riego_cat_ha=tot['sin_riego_catastral_ha'])


def _tabla_universos(u):
    return [
        ['Fichas', esn(u['principales'] + u['adicionales']),
         f"{esn(u['principales'])} principales (una por titular entrevistado) y {esn(u['adicionales'])} adicionales "
         '(otros predios del mismo titular).'],
        ['Titulares únicos', esn(u['cedulas']),
         f"Cédulas distintas entre las fichas principales. {esn(u['sin_cedula'])} fichas principales no registran "
         'cédula; la verificación de cédulas está en curso.'],
        ['Predios', esn(u['predios']),
         f"Predios del catastro del cantón con al menos una ficha ({esn(u['cat_ha'], 2)} ha catastrales). Un predio "
         'puede tener varias fichas.'],
        ['Fichas sin polígono', esn(u['sinpol']),
         'Fichas cuyo predio no está en el catastro del cantón (ver tabla fichas_sin_poligono).'],
        ['Comunidades', esn(u['comunidades']), 'Comunidades del catálogo con fichas.'],
    ]


def _tabla_superficies(u):
    return [
        ['Catastral', esn(u['cat_ha'], 2), 'area_catastro_m2', 'El polígono del catastro municipal; cada predio '
         'cuenta una vez.'],
        ['Con riego (catastral)', esn(u['riego_ajustado_ha'], 2), '—', 'Parte de la catastral bajo riego, ajustada '
         'donde lo declarado excede el predio.'],
        ['Sin riego (catastral)', esn(u['sin_riego_cat_ha'], 2), '—', 'Catastral menos con riego.'],
        ['Declarada', esn(u['decl_ha'], 2), 'area_total_m2', 'Lo que informó cada titular, sumado ficha por ficha.'],
        ['Declarada con riego', esn(u['riego_ha'], 2), 'area_riego_m2', 'Parte de la declarada bajo riego.'],
        ['Declarada sin riego', esn(u['sin_riego_ha'], 2), 'area_sin_riego_m2', 'Parte de la declarada sin riego.'],
        ['Declarada en predios con polígono', esn(u['decl_pol_ha'], 2), 'area_declarada_m2',
         'La declarada sin las fichas sin polígono.'],
    ]


def diccionario_v2(gpkg, shp_dir, salida_dir):
    est = estructura_v2(gpkg, shp_dir)
    u = universos(gpkg)
    os.makedirs(salida_dir, exist_ok=True)
    # Excel
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from excel_compat import aplicar_formatos
    xls = os.path.join(salida_dir, 'DICCIONARIO DE DATOS v2.xlsx')
    wb = Workbook(); ws = wb.active; ws.title = 'Diccionario'
    ws.append(['Capa', 'Geometría', 'Registros', 'Campo', 'En el Shapefile', 'Tipo', 'Descripción',
               'Dominio / valores admitidos', 'Vacíos'])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1E3A8A')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for capa in est:
        for i, f in enumerate(capa['campos']):
            ws.append([capa['capa'] if i == 0 else '', capa['geometria'] if i == 0 else '',
                       capa['registros'] if i == 0 else '', f['campo'], f['shp'], f['tipo'], f['descripcion'],
                       f['dominio'], f['vacios']])
    for col, ancho in zip('ABCDEFGHI', [24, 11, 10, 24, 13, 14, 70, 40, 14]):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for nombre, cab, filas in (('Universos', ['Universo', 'Cantidad', 'Definición'], _tabla_universos(u)),
                               ('Superficies', ['Superficie', 'Hectáreas', 'Campo', 'Definición'],
                                _tabla_superficies(u))):
        w = wb.create_sheet(nombre); w.append(cab)
        for c in w[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1E3A8A')
        for f in filas:
            w.append(f)
        for col, ancho in zip('ABCD', [32, 14, 22, 80]):
            w.column_dimensions[col].width = ancho
    wb.save(xls); aplicar_formatos(xls)
    # PDF
    pdf = os.path.join(salida_dir, 'DICCIONARIO DE DATOS v2.pdf')
    st = ap5._estilos()
    h = [Paragraph('Diccionario de datos · versión 2', st['h']), Spacer(0, 1.5 * mm),
         Paragraph('Catastro socioeconómico y productivo predial · Sistema de riego comunitario Guanguilquí–Porotog',
                   st['sub']), Spacer(0, 5 * mm),
         Paragraph('Describe la estructura, los campos, los dominios y el contenido de la base de datos geoespacial '
                   'del catastro, en su versión 2. La estructura, los dominios y los vacíos se leen del propio '
                   'GeoPackage, no se escriben a mano. Los datos corresponden al levantamiento de campo cerrado al '
                   f'{FECHA_CORTE}; la versión 2 no cambia ninguna cifra de la versión 1: agrega campos, tablas y '
                   'definiciones.', st['p'])]
    h += titulo_seccion('Universos: fichas, titulares y predios no son lo mismo')
    h.append(tabla_datos(['Universo', 'Cantidad', 'Definición'], _tabla_universos(u), [38 * mm, 22 * mm, 130 * mm]))
    h += [Spacer(0, 2 * mm), Paragraph('Una ficha no es un usuario ni un predio: un titular puede tener varias fichas '
                                       '(sus predios) y un predio puede tener varias fichas (los herederos de un '
                                       'terreno familiar). Por eso cada cifra se lee dentro de su universo.', st['p'])]
    h += titulo_seccion('Superficies: dos mediciones que no se mezclan')
    h.append(tabla_datos(['Superficie', 'Hectáreas', 'Campo', 'Definición'], _tabla_superficies(u),
                         [45 * mm, 22 * mm, 33 * mm, 90 * mm]))
    h += [Spacer(0, 2 * mm), Paragraph('Las superficies se guardan en metros cuadrados con dos decimales y se '
                                       'redondean a hectáreas solo al presentarlas. Ninguna de ellas es área '
                                       'potencialmente regable: esa definición corresponde al estudio agrícola.',
                                       st['p'])]
    h += titulo_seccion('Las capas y tablas')
    h.append(tabla_datos(['Capa', 'Geometría', 'Registros', 'Qué contiene'],
                         [[c['capa'], c['geometria'], esn(c['registros']), c['proposito']] for c in est],
                         [42 * mm, 18 * mm, 20 * mm, 110 * mm]))
    h += titulo_seccion('Cómo se enlazan')
    h += [Paragraph('<b>clave_catastral</b> · enlaza predios_investigados, catastro_completo, fichas y '
                    'correspondencia_fichas.', st['li']),
          Paragraph('<b>codigo_ficha</b> · identificador único de cada ficha (S01-C22-R001-F01). Está en fichas, en '
                    'correspondencia_fichas y, agrupado, en codigos_ficha de predios_investigados; es el mismo con el '
                    'que se nombra cada ficha en PDF.', st['li']),
          Paragraph('<b>ficha_id</b> · enlaza cada ficha con sus cultivos, sus animales y su registro de cambios.',
                    st['li']),
          Paragraph('<b>codigo_comunidad</b> · enlaza las capas con el catálogo de comunidades.', st['li']),
          Spacer(0, 2 * mm),
          Paragraph('<b>Nombres en el Shapefile.</b> El formato Shapefile admite nombres de hasta 10 caracteres. La '
                    'columna «Shapefile» de cada tabla da el nombre que tiene cada campo en ese formato; algunos se '
                    'abreviaron para que se entiendan (TITULARES, PROP_CATAS, TEL_CEL, TEL_CASA, SIST_RIEGO, COM_OFIC, '
                    'COD_PRINC). Al abrir un Shapefile en una hoja de cálculo, usar la versión Excel de las tablas: '
                    'el Shapefile está en UTF-8 y una hoja de cálculo que no lo lea así altera las tildes, y las '
                    'claves catastrales largas solo se conservan completas como texto.', st['p'])]
    for capa in est:
        h += titulo_seccion(f"Capa: {capa['capa']}")
        h += [Paragraph(f"{capa['geometria']} · {esn(capa['registros'])} registros. {capa['proposito']}", st['p']),
              Spacer(0, 2 * mm)]
        h.append(tabla_datos(['Campo', 'Shapefile', 'Tipo', 'Descripción', 'Dominio', 'Vacíos'],
                             [[f['campo'], f['shp'], f['tipo'], f['descripcion'], f['dominio'] or '—', f['vacios']]
                              for f in capa['campos']],
                             [30 * mm, 19 * mm, 17 * mm, 74 * mm, 32 * mm, 18 * mm]))
    ap5._cerrar(ap5._doc(pdf, 'Diccionario de datos · versión 2'), h)
    sin_desc = [(c['capa'], f['campo']) for c in est for f in c['campos'] if not f['descripcion']]
    return pdf, xls, sin_desc


def nota_avance(gpkg, carpeta):
    u = universos(gpkg)
    ruta = os.path.join(carpeta, 'NOTA ACLARATORIA - porcentaje de avance (observacion C05).pdf')
    st = ap5._estilos()
    h = [Paragraph('Nota aclaratoria al documento «Porcentaje de avance y fecha de corte»', st['h']),
         Spacer(0, 1.5 * mm),
         Paragraph('Producto 5 · Respuesta a la observación C05 del informe JAFG-2026-002', st['sub']),
         Spacer(0, 5 * mm),
         Paragraph('Esta nota complementa el documento 6 de la entrega del 18 de septiembre de 2026, que se mantiene '
                   'sin cambios. La observación C05 señala que el avance del 100 % se declara sin un universo de '
                   'predios conciliado. La nota precisa sobre qué universo se mide cada cifra.', st['p'])]
    h += titulo_seccion('Qué afirma el documento 6')
    h += [Paragraph('El 100 % se refiere a las propiedades declaradas por los comuneros que se presentaron al '
                    'levantamiento: todas las que declararon tienen su ficha. El mismo documento advierte que no es '
                    'posible un comparativo exacto contra el número total de propiedades, porque no se dispone de '
                    'cuántas propiedades tienen en total los comuneros de los listados.', st['p'])]
    h += titulo_seccion('El avance por universo')
    h.append(tabla_datos(['Universo', 'Cantidad', 'Definición'], _tabla_universos(u), [38 * mm, 22 * mm, 130 * mm]))
    h += [Spacer(0, 2 * mm),
          Paragraph('Como referencia, los listados de las 50 comunidades registran 3.681 comuneros. Las fichas no se '
                    'comparan con ese número: una ficha no es un usuario.', st['p'])]
    h += titulo_seccion('Lo que requiere información de terceros')
    h += [Paragraph('El porcentaje sobre el total de predios de los usuarios del sistema solo puede calcularse con '
                    'los padrones de usuarios certificados por cada organización, que indiquen quiénes son sus '
                    'usuarios y qué predios tiene cada uno (punto 11.1 del informe). La Consultora no dispone de '
                    'esos padrones.', st['p']),
          Spacer(0, 2 * mm),
          Paragraph(f'La fecha de corte de la información se mantiene: {FECHA_CORTE}.', st['p'])]
    ap5._cerrar(ap5._doc(ruta, 'Nota aclaratoria al porcentaje de avance'), h)
    return ruta


def guia_v2(destino, cambios):
    ruta = os.path.join(destino, '0 - QUE CAMBIA EN LA VERSION 2.pdf')
    st = ap5._estilos()
    h = [Paragraph('Qué cambia en la versión 2', st['h']), Spacer(0, 1.5 * mm),
         Paragraph('Producto 5 · Reentrega con los ajustes de la respuesta al informe JAFG-2026-002', st['sub']),
         Spacer(0, 5 * mm),
         Paragraph('Esta versión parte de la base entregada el 18 de septiembre de 2026 y solo agrega: campos, tablas '
                   'y documentos. Ninguna cifra cambia: se comparó capa por capa con la versión 1 y coinciden el '
                   'número de registros, las geometrías, las superficies y los conteos. La carpeta conserva la misma '
                   'numeración; el archivo «0 - INVENTARIO DE ARCHIVOS» dice cuáles son nuevos, cuáles se '
                   'actualizaron y cuáles no cambiaron.', st['p'])]
    h += titulo_seccion('Lo que se agregó y a qué observación responde')
    filas = [
        ['Códigos de ficha en la capa de predios (codigos_ficha)', 'C04, C10'],
        ['Código y nombre oficial de comunidad, y sector 1–3, en las capas; catálogo de las 50 comunidades', 'C07'],
        ['«sector_riego» pasa a «sistema_riego_declarado», que es lo que contiene', '7.1-c'],
        ['Fecha de corte y versión en las capas y en la tabla metadatos', '7.1-b'],
        ['Tabla de correspondencia ficha – titular – predio – polígono, con marcas de revisión', 'C04, 11.1-c'],
        ['Tabla de las 9 fichas sin polígono y su causa', 'C09'],
        ['Cultivos: donde se llenó el tipo y el campo «otro», manda el tipo (el «500» es «Pasto no mejorado»)',
         'C17'],
        ['Registro de cambios ficha por ficha desde el 18 de agosto de 2026', '11.1-c'],
        ['Tablas en Excel con las claves catastrales como texto (carpeta 1.4)', 'C10, C17'],
        ['Shapefile regenerado con todos los campos, incluidas las superficies', '7.1-a'],
        ['Diccionario de datos v2: universos, superficies, dominios, vacíos y nombres en el Shapefile',
         'C03, C17, 9-a, 7.1-b'],
        ['Nota aclaratoria al porcentaje de avance (carpeta 6)', 'C05'],
    ]
    h.append(tabla_datos(['Ajuste', 'Observación'], filas, [160 * mm, 30 * mm]))
    h += titulo_seccion('Lo que no cambia')
    h += [Paragraph('Las fichas, los predios, las geometrías, las superficies, los cultivos y los animales son los '
                    'mismos de la versión 1. No se eliminó ni se fusionó ningún registro: los posibles registros '
                    'duplicados y las cédulas en verificación quedan marcados en la tabla de correspondencia. El '
                    'formato CAD, la memoria técnica, el manual y el enlace a la plataforma no cambian.', st['p'])]
    h += titulo_seccion('Registro de cambios')
    h += [Paragraph(f'Desde el 18 de agosto de 2026 la base registra {cambios}. El detalle, ficha por ficha y con '
                    'el valor anterior y el nuevo, está en la tabla registro_cambios; los valores originales se '
                    'conservan en los respaldos de la base.', st['p'])]
    ap5._cerrar(ap5._doc(ruta, 'Qué cambia en la versión 2'), h)
    return ruta


DESCRIPCION_ARCHIVO = [
    ('QUE CAMBIA EN LA VERSION 2', 'Guía de la versión 2', 'Todas'),
    ('INVENTARIO DE ARCHIVOS', 'Este inventario', '7.1-a, 11.1-b'),
    ('LEA ESTO PRIMERO', 'Guía de la entrega v1 y su correspondencia con el contrato', '—'),
    ('padron_riego_porotog.gpkg', 'Base geoespacial completa (GeoPackage)', 'C04, C07, C09, C10, C17, 7.1, 11.1-c'),
    ('padron_riego_porotog.qgz', 'Proyecto QGIS simbolizado', '7.1-a'),
    ('TABLAS DE ATRIBUTOS', 'Tablas en Excel, con claves como texto', 'C10, C17'),
    ('DICCIONARIO DE DATOS', 'Diccionario de datos', 'C03, C17, 9-a, 7.1-b'),
    ('NOTA ACLARATORIA', 'Nota aclaratoria al porcentaje de avance', 'C05'),
    ('PORCENTAJE DE AVANCE', 'Porcentaje de avance y fecha de corte (v1)', '—'),
    ('LINK DE LA PLATAFORMA', 'Enlace a la plataforma', '—'),
    ('MEMORIA TECNICA', 'Memoria técnica del geovisor', '—'),
    ('MANUAL DE USO', 'Manual de uso del geovisor', '—'),
    ('.dxf', 'Capa en formato CAD', '7.1-a'),
    ('.shp', 'Capa en formato Shapefile', '7.1-a'),
    ('.dbf', 'Tabla del Shapefile', '7.1-a'),
]


def inventario(destino, zip_v1):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from excel_compat import aplicar_formatos
    v1 = {}
    with zipfile.ZipFile(zip_v1) as z:
        raiz = z.namelist()[0].split('/')[0]
        for i in z.infolist():
            if not i.filename.endswith('/'):
                v1[i.filename[len(raiz) + 1:].replace('/', os.sep)] = hashlib.md5(z.read(i)).hexdigest()
    ruta = os.path.join(destino, '0 - INVENTARIO DE ARCHIVOS.xlsx')
    filas = []
    for r, _, fs in os.walk(destino):
        for f in sorted(fs):
            p = os.path.join(r, f); rel = os.path.relpath(p, destino)
            if rel == os.path.basename(ruta):
                continue
            h = hashlib.md5(open(p, 'rb').read()).hexdigest()
            estado = 'Nuevo en la v2' if rel not in v1 else ('Sin cambios' if v1[rel] == h else 'Actualizado en la v2')
            desc, obs = next(((d, o) for k, d, o in DESCRIPCION_ARCHIVO if k.lower() in f.lower()), ('', ''))
            if f.lower().endswith(('.shx', '.prj', '.cpg')):
                desc, obs = 'Archivo auxiliar del Shapefile', '7.1-a'
            filas.append((os.path.dirname(rel) or '(raíz)', f, round(os.path.getsize(p) / 1024, 1), estado, desc, obs))
    retirados = sorted(k for k in v1 if not os.path.exists(os.path.join(destino, k)))
    wb = Workbook(); ws = wb.active; ws.title = 'Inventario'
    ws.append(['Carpeta', 'Archivo', 'Tamaño (KB)', 'Estado frente a la v1', 'Qué es', 'Observación que atiende'])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1E3A8A')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for f in filas:
        ws.append(list(f))
    for k in retirados:
        ws.append([os.path.dirname(k), os.path.basename(k), '', 'Sustituido en la v2',
                   'Lo reemplaza su versión 2 (ver «Qué cambia en la versión 2»)', ''])
    for col, ancho in zip('ABCDEF', [52, 46, 12, 22, 46, 26]):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical='top', wrap_text=True)
        fila[2].number_format = '#,##0.0'
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    wb.save(ruta); aplicar_formatos(ruta)
    return ruta, collections.Counter(f[3] for f in filas), retirados
