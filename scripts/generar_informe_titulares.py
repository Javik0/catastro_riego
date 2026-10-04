# -*- coding: utf-8 -*-
"""
Informe de sustento: titulares entrevistados y conocimiento del proyecto.

Es el documento que acompaña al listado en Excel
(`generar_listado_titulares.py`) y se entrega al sociólogo (Alexis Guerrero),
que necesita sostener cuántos titulares fueron entrevistados y cuántos
contestaron que conocen el proyecto.

Tono y alcance
--------------
Neutro: expone lo que consta en las fichas y dice sobre qué base se calcula
cada porcentaje. No sale ninguna cifra escrita a mano: todo se cuenta del
mismo padrón y con la misma función de carga que el Excel, para que los dos
archivos no puedan contradecirse.

Solo cita como evidencia lo que el dato tiene: fecha, técnico, coordenadas,
clave catastral y cédula. El campo de consentimiento informado está vacío en
las 4.307 fichas y por eso el documento no lo menciona como respaldo.

Uso:  python -X utf8 scripts/generar_informe_titulares.py
"""
import collections
import os
import sys
from datetime import datetime

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

from generar_listado_titulares import cargar  # noqa: E402
from informe_estilo import esn  # noqa: E402
from generar_fichas_pdf import (  # noqa: E402
    AZUL as AZUL_PDF, TINTA, FECHA_CORTE, PROJECT_LOCATION, PROJECT_SUBTITLE,
    PROJECT_SUBTITLE_ALCANCE, PROJECT_TITLE, PUB, colors, mm, Paragraph,
    ParagraphStyle, Spacer, tabla_datos, titulo_seccion,
)
import a_word  # noqa: E402

SALIDA = (r'C:\Users\HP\OneDrive\Escritorio'
          r'\INFORME DE TITULARES ENTREVISTADOS - conocimiento del proyecto.docx')

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
         'septiembre', 'octubre', 'noviembre', 'diciembre']


def fecha_larga(d):
    return f'{d.day} de {MESES[d.month - 1]} de {d.year}'


def pc(a, b):
    """Porcentaje en castellano (coma decimal) o raya si la base es cero."""
    return f'{esn(100 * a / b, 1)} %' if b else '—'


def main():
    filas, por_cedula = cargar()
    N = len(filas)
    si = sum(1 for x in filas if x['conoce'] == 'Sí')
    no = sum(1 for x in filas if x['conoce'] == 'No')
    sr = N - si - no
    resp = si + no

    fechas = [datetime.strptime(x['fecha'], '%d/%m/%Y') for x in filas if x['fecha']]
    n_comunidades = len({x['com_n'] for x in filas})
    n_tecnicos = len({x['tecnico'] for x in filas if x['tecnico']})
    n_dias = len({d.date() for d in fechas})
    con_fecha = len(fechas)
    con_tec = sum(1 for x in filas if x['tecnico'])
    con_gps = sum(1 for x in filas if x['gps'])
    con_clave = sum(1 for x in filas if x['clave'])
    con_nombre = sum(1 for x in filas if x['nombre'])

    ced = collections.Counter(x['estado_ced'] for x in filas)
    sin_ced = ced['Sin cédula']
    con_ced = N - sin_ced
    dup = {c: v for c, v in por_cedula.items() if len(v) > 1}
    en_dup = sum(len(v) for v in dup.values())

    # Cortes por sector y por comunidad (misma regla que el Excel: el sector y
    # la comunidad salen del código de la ficha).
    por_sector = collections.OrderedDict()
    por_com = collections.OrderedDict()
    for x in sorted(filas, key=lambda x: (x['com_n'], x['codigo'])):
        for d in (por_sector.setdefault(x['sector'], collections.Counter()),
                  por_com.setdefault((x['com_n'], x['comunidad'], x['sector']),
                                     collections.Counter())):
            d['n'] += 1
            d['si' if x['conoce'] == 'Sí' else 'no' if x['conoce'] == 'No' else 'sr'] += 1
    assert sum(d['n'] for d in por_sector.values()) == N
    assert sum(d['n'] for d in por_com.values()) == N
    assert sum(d['si'] for d in por_com.values()) == si

    st_h = ParagraphStyle('h', fontName='Helvetica-Bold', fontSize=13, leading=16,
                          textColor=AZUL_PDF, alignment=1)
    st_sub = ParagraphStyle('sub', fontName='Helvetica', fontSize=9.5, leading=12,
                            alignment=1, textColor=colors.HexColor('#475569'))
    st_p = ParagraphStyle('p', fontName='Helvetica', fontSize=9.5, leading=13,
                          textColor=TINTA, alignment=4)
    st_li = ParagraphStyle('li', parent=st_p, leftIndent=6 * mm, spaceAfter=2)

    b = [Paragraph('Titulares entrevistados y conocimiento del proyecto', st_h),
         Spacer(0, 1.5 * mm),
         Paragraph(f'Informe de sustento · Padrón de usuarios · Información al {FECHA_CORTE}',
                   st_sub),
         Spacer(0, 5 * mm),
         Paragraph('Este documento resume el levantamiento del padrón de usuarios del sistema '
                   'de riego comunitario Guanguilquí–Porotog: cuántas fichas principales se '
                   'levantaron, cuántas respondieron la pregunta «¿Conoce el proyecto?» y qué '
                   'consta en cada ficha. Acompaña al listado en Excel, que detalla cada caso '
                   'por comunidad.', st_p)]

    # 1 ─ El levantamiento
    b += titulo_seccion('1. El levantamiento')
    b.append(Paragraph(
        f'El levantamiento se hizo casa por casa entre el {fecha_larga(min(fechas))} y el '
        f'{fecha_larga(max(fechas))}, en {esn(n_comunidades)} comunidades de los tres sectores, '
        f'con {esn(n_tecnicos)} técnicos investigadores y {esn(n_dias)} días con registros. '
        f'Cada ficha principal recoge la entrevista a un titular; las fichas adicionales son '
        f'otros predios y no se cuentan como titulares. En total hay '
        f'<b>{esn(N)} fichas principales</b>.', st_p))

    # 2 ─ Respuesta a la pregunta
    b += titulo_seccion('2. Respuesta a «¿Conoce el proyecto?»')
    b.append(tabla_datos(
        ['Concepto', 'Fichas', '% de las fichas', '% de quienes respondieron'],
        [['Fichas principales', esn(N), '100,0 %', '—'],
         ['Respondieron la pregunta', esn(resp), pc(resp, N), '100,0 %'],
         ['Conocen el proyecto (Sí)', esn(si), pc(si, N), pc(si, resp)],
         ['No conocen el proyecto (No)', esn(no), pc(no, N), pc(no, resp)],
         ['Sin respuesta registrada', esn(sr), pc(sr, N), '—']],
        [70 * mm, 24 * mm, 32 * mm, 44 * mm]))
    b.append(Spacer(0, 2 * mm))
    b.append(Paragraph(
        f'La pregunta fue respondida en {esn(resp)} fichas ({pc(resp, N)}). Entre quienes '
        f'respondieron, {esn(si)} ({pc(si, resp)}) dicen que conocen el proyecto. Si la '
        f'proporción se calcula sobre las {esn(N)} fichas principales, es {pc(si, N)}. Ambas '
        f'lecturas parten de las mismas cifras; solo cambia la base del cálculo.', st_p))

    # 3 ─ Por sector
    b += titulo_seccion('3. Por sector')
    filas_s = [[sec, esn(d['n']), esn(d['si'] + d['no']), esn(d['si']), esn(d['no']),
                esn(d['sr']), pc(d['si'], d['si'] + d['no'])]
               for sec, d in sorted(por_sector.items())]
    filas_s.append(['Total', esn(N), esn(resp), esn(si), esn(no), esn(sr), pc(si, resp)])
    b.append(tabla_datos(
        ['Sector', 'Fichas', 'Respondieron', 'Sí', 'No', 'Sin respuesta',
         '% Sí (sobre quienes respondieron)'],
        filas_s, [26 * mm, 20 * mm, 26 * mm, 18 * mm, 18 * mm, 24 * mm, 38 * mm],
        ultima_negrita=True))

    # 4 ─ Qué consta en cada ficha
    b += titulo_seccion('4. Qué consta en cada ficha')
    b.append(tabla_datos(
        ['Dato registrado', 'Fichas', '% de las fichas'],
        [['Apellidos y nombres del titular', esn(con_nombre), pc(con_nombre, N)],
         ['Clave catastral del predio', esn(con_clave), pc(con_clave, N)],
         ['Fecha de levantamiento', esn(con_fecha), pc(con_fecha, N)],
         ['Técnico investigador que la levantó', esn(con_tec), pc(con_tec, N)],
         ['Coordenadas del punto de levantamiento', esn(con_gps), pc(con_gps, N)]],
        [90 * mm, 30 * mm, 40 * mm]))

    # 5 ─ Identificación de los titulares
    b += titulo_seccion('5. Identificación de los titulares')
    b.append(tabla_datos(
        ['Cédula', 'Fichas', '% de las fichas'],
        [['Con cédula registrada', esn(con_ced), pc(con_ced, N)],
         ['   con dígito verificador válido', esn(ced['Dígito verificador válido']),
          pc(ced['Dígito verificador válido'], N)],
         ['   con dígito verificador no válido', esn(ced['Dígito verificador NO válido']),
          pc(ced['Dígito verificador NO válido'], N)],
         ['   en otro formato (RUC o pasaporte)', esn(ced['Otro formato (RUC / pasaporte)']),
          pc(ced['Otro formato (RUC / pasaporte)'], N)],
         ['Sin cédula registrada', esn(sin_ced), pc(sin_ced, N)]],
        [90 * mm, 30 * mm, 40 * mm]))
    b.append(Spacer(0, 2 * mm))
    b.append(Paragraph(
        f'Las {esn(N)} fichas principales corresponden a {esn(len(por_cedula))} cédulas '
        f'distintas y {esn(sin_ced)} fichas sin cédula registrada. {esn(len(dup))} cédulas '
        f'figuran en más de una ficha principal ({esn(en_dup)} fichas). El dígito verificador '
        f'comprueba que el número de cédula está bien formado; la consulta al Registro Civil '
        f'no forma parte de este levantamiento.', st_p))

    # 6 ─ Archivos que acompañan
    b += titulo_seccion('6. Listado que acompaña a este informe')
    b.append(Paragraph('El Excel «Listado de titulares entrevistados» trae estas hojas:', st_p))
    b.append(Spacer(0, 1 * mm))
    for t in [
        f'<b>LISTADO COMPLETO:</b> las {esn(N)} fichas principales, con código de ficha, '
        'sector, comunidad, cédula, apellidos y nombres, respuesta, clave catastral, fecha de '
        'levantamiento y técnico.',
        f'<b>CONOCEN EL PROYECTO:</b> las {esn(si)} fichas que respondieron que sí, con las '
        'mismas columnas.',
        '<b>POR COMUNIDAD:</b> fichas, respuestas y porcentajes de cada comunidad.',
        '<b>RESUMEN:</b> las cifras de este informe.',
    ]:
        b.append(Paragraph(t, st_li))

    # Anexo ─ por comunidad
    b += titulo_seccion('Anexo. Respuesta por comunidad')
    filas_c = [[nom, sec, esn(d['n']), esn(d['si'] + d['no']), esn(d['si']),
                pc(d['si'], d['si'] + d['no'])]
               for (_, nom, sec), d in por_com.items()]
    filas_c.append(['Total', '', esn(N), esn(resp), esn(si), pc(si, resp)])
    b.append(tabla_datos(
        ['Comunidad', 'Sector', 'Fichas', 'Respondieron', 'Sí conocen',
         '% Sí (sobre quienes respondieron)'],
        filas_c, [58 * mm, 20 * mm, 18 * mm, 26 * mm, 22 * mm, 34 * mm],
        ultima_negrita=True))

    a_word.escribir(b, SALIDA, cabecera_estudio=dict(
        titulo=PROJECT_TITLE, subtitulo=PROJECT_SUBTITLE,
        alcance=PROJECT_SUBTITLE_ALCANCE, ubicacion=PROJECT_LOCATION,
        logo_izq=os.path.join(PUB, 'logo-izq.png'),
        logo_der=os.path.join(PUB, 'logo-der.png'),
        pie_izquierda='AP&CATASTROS',
        pie_derecha_prefijo=f'CONSORCIO CAYAMBE SPT · Datos al {FECHA_CORTE} · pág. '))

    # Se vuelve a leer el .docx: lo que se entrega es lo que quedó en disco.
    import docx
    d = docx.Document(SALIDA)
    texto = '\n'.join(p.text for p in d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            texto += '\n' + ' | '.join(c.text for c in r.cells)
    for esperado in (esn(N), esn(si), pc(si, resp), pc(si, N), esn(len(por_cedula))):
        assert esperado in texto, f'falta «{esperado}» en el Word'
    assert len(d.tables) == 5, len(d.tables)
    assert len(d.tables[-1].rows) == n_comunidades + 2, len(d.tables[-1].rows)

    print(f'Fichas principales : {esn(N)}')
    print(f'  Sí / respondieron: {esn(si)} / {esn(resp)} = {pc(si, resp)} '
          f'(sobre fichas: {pc(si, N)})')
    print(f'  cédulas distintas: {esn(len(por_cedula))} · sin cédula {esn(sin_ced)}')
    print(f'  levantamiento    : {fecha_larga(min(fechas))} a {fecha_larga(max(fechas))} · '
          f'{n_tecnicos} técnicos · {n_dias} días')
    print('Contenido releído del .docx: cifras y tablas presentes.')
    print('✔', SALIDA)
    return 0


if __name__ == '__main__':
    sys.exit(main())
