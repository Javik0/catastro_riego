# -*- coding: utf-8 -*-
"""
Convierte a Word los mismos bloques con que se arma un PDF de reportlab.

Por que asi y no reescribiendo el texto: si el documento se redactara dos
veces —una para el PDF y otra para el Word— con el tiempo se separarian y
nadie sabria cual manda. Aqui hay una sola fuente: los generadores construyen
su lista de bloques como siempre, y este modulo la pinta en Word.

La plantilla es el propio `membrete.docx` de la Prefectura, de modo que el
Word sale con su cabecera y su pie reales, no con una imitacion.
"""
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AZUL = RGBColor(0x1E, 0x3A, 0x8A)
GRIS = RGBColor(0x47, 0x55, 0x69)
TINTA = RGBColor(0x1F, 0x29, 0x37)

ALINEACION = {
    0: WD_ALIGN_PARAGRAPH.LEFT,
    1: WD_ALIGN_PARAGRAPH.CENTER,
    2: WD_ALIGN_PARAGRAPH.RIGHT,
    4: WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _limpiar(t):
    """El texto de reportlab trae marcado propio; se pasa a texto llano."""
    t = re.sub(r'<br\s*/?>', '\n', t or '')
    t = t.replace('&nbsp;', ' ').replace('&amp;', '&')
    t = t.replace('&lt;', '<').replace('&gt;', '>')
    return t


def _escribir_con_negritas(parrafo, texto, tam, color):
    """Reparte el texto en tramos, respetando los <b> del original."""
    for tramo in re.split(r'(<b>.*?</b>)', _limpiar(texto)):
        if not tramo:
            continue
        negrita = tramo.startswith('<b>')
        limpio = re.sub(r'</?b>', '', tramo)
        limpio = re.sub(r'<[^>]+>', '', limpio)      # cualquier otra etiqueta
        if not limpio:
            continue
        r = parrafo.add_run(limpio)
        r.bold = negrita
        r.font.size = Pt(tam)
        r.font.color.rgb = color
        r.font.name = 'Calibri'


def _es_titulo_seccion(tabla):
    """titulo_seccion() devuelve una tabla de una sola celda con el rotulo."""
    return len(tabla._cellvalues) == 1 and len(tabla._cellvalues[0]) == 1


def _texto_celda(c):
    """El texto de una celda, venga como venga.

    Algunas celdas no traen el parrafo suelto sino dentro de una lista —es lo
    que hace titulo_seccion()—, y sin desenvolverlo se acaba escribiendo el
    objeto entero en el documento.
    """
    if c is None:
        return ''
    if isinstance(c, (list, tuple)):
        return '\n'.join(_texto_celda(x) for x in c if x is not None).strip()
    if hasattr(c, 'text'):
        return re.sub(r'<[^>]+>', '', _limpiar(c.text))
    return str(c)


def _poner_tabla(doc, tabla):
    filas = tabla._cellvalues
    if not filas:
        return
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, celda in enumerate(fila):
            if j >= len(t.columns):
                continue
            p = t.cell(i, j).paragraphs[0]
            r = p.add_run(_texto_celda(celda))
            r.font.size = Pt(8.5)
            r.font.name = 'Calibri'
            if i == 0:                      # la primera fila es la cabecera
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _sombrear(t.cell(i, j), '1E3A8A')
    doc.add_paragraph()


def _sombrear(celda, hex_color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    sombra = OxmlElement('w:shd')
    sombra.set(qn('w:val'), 'clear')
    sombra.set(qn('w:fill'), hex_color)
    celda._tc.get_or_add_tcPr().append(sombra)


def _pintar(doc, bloques, primero):
    """Recorre los bloques del PDF y los va escribiendo en el documento."""
    for b in bloques:
        nombre = type(b).__name__

        if nombre == 'KeepTogether':
            _pintar(doc, getattr(b, '_content', []), primero)
            continue

        if nombre == 'Spacer':
            continue

        if nombre == 'Image':
            # reportlab guarda `str(buf)` en .filename cuando la imagen viene
            # de un BytesIO —los graficos de matplotlib se pasan asi, sin
            # tocar disco—, y el buffer real se pierde. Por eso los
            # generadores dejan los bytes aparte, en `_png_bytes` (ver
            # `_imagen()` en el script del informe). Sin eso, se intenta la
            # ruta de archivo de siempre.
            datos = getattr(b, '_png_bytes', None)
            origen = io.BytesIO(datos) if datos else getattr(b, 'filename', None)
            usable = isinstance(origen, io.BytesIO) or (
                isinstance(origen, str) and os.path.exists(origen))
            if usable:
                from docx.shared import Mm
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # El ancho util de la hoja son unos 165 mm.
                ancho = min(165, (getattr(b, 'drawWidth', 400) or 400) / 72 * 25.4)
                p.add_run().add_picture(origen, width=Mm(ancho))
            continue

        if nombre == 'Table':
            if _es_titulo_seccion(b):
                rotulo = _texto_celda(b._cellvalues[0][0])
                p = doc.add_paragraph()
                p.space_before = Pt(10)
                r = p.add_run(rotulo)
                r.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = AZUL
                r.font.name = 'Calibri'
            else:
                _poner_tabla(doc, b)
            continue

        if nombre == 'Paragraph':
            estilo = getattr(b, 'style', None)
            tam = getattr(estilo, 'fontSize', 9) or 9
            alin = getattr(estilo, 'alignment', 0)
            sangria = getattr(estilo, 'leftIndent', 0) or 0
            p = doc.add_paragraph()
            p.alignment = ALINEACION.get(alin, WD_ALIGN_PARAGRAPH.LEFT)
            if sangria:
                from docx.shared import Mm
                p.paragraph_format.left_indent = Mm(sangria / 72 * 25.4)
            color = TINTA
            if alin == 1 and tam >= 12:
                color, tam = AZUL, 14          # titulo del documento
            elif alin == 1 and tam < 10:
                color = GRIS                   # subtitulo
            _escribir_con_negritas(p, getattr(b, 'text', ''), tam, color)
            continue
    return doc


def _quitar_bordes_tabla(tabla):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    bordes = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'none')
        bordes.append(el)
    tabla._tbl.tblPr.append(bordes)


def agregar_cabecera_pie_estudio(doc, *, titulo, subtitulo, alcance='', ubicacion='',
                                 logo_izq=None, logo_der=None,
                                 pie_izquierda='', pie_derecha_prefijo=''):
    """La cabecera y el pie del estudio, cuando el documento NO lleva el
    membrete de la Prefectura.

    No hay una plantilla .docx para esta cabecera —a diferencia del membrete,
    que sí es un archivo real que entregaron—, así que se arma con una tabla
    de tres columnas (logo · títulos centrados · logo) en el header y un
    número de página real en el footer, replicando lo que dibuja
    `cabecera_pie()` en el PDF del mismo informe.
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AP
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Mm, Pt

    sec = doc.sections[0]
    sec.header_distance = Mm(8)
    sec.footer_distance = Mm(8)

    header = sec.header
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)
    t = header.add_table(rows=1, cols=3, width=Mm(170))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Mm(28)
    t.columns[1].width = Mm(114)
    t.columns[2].width = Mm(28)
    _quitar_bordes_tabla(t)

    if logo_izq and os.path.exists(logo_izq):
        t.rows[0].cells[0].paragraphs[0].add_run().add_picture(logo_izq, width=Mm(22))

    celda = t.rows[0].cells[1]
    celda.paragraphs[0].alignment = AP.CENTER
    lineas = [(titulo, 6.3, GRIS, True), (subtitulo, 9.5, TINTA, True),
              (alcance, 6.3, GRIS, False), (ubicacion, 6.3, GRIS, False)]
    primero = True
    for texto, tam, color, negrita in lineas:
        if not texto:
            continue
        p = celda.paragraphs[0] if primero else celda.add_paragraph()
        p.alignment = AP.CENTER
        primero = False
        r = p.add_run(texto)
        r.font.size = Pt(tam)
        r.font.color.rgb = color
        r.font.bold = negrita
        r.font.name = 'Calibri'

    if logo_der and os.path.exists(logo_der):
        p = t.rows[0].cells[2].paragraphs[0]
        p.alignment = AP.RIGHT
        p.add_run().add_picture(logo_der, width=Mm(22))

    # La raya azul bajo la cabecera, como en el PDF.
    borde = header.add_paragraph()
    pPr = borde._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '18')
    bottom.set(qn('w:color'), '1E3A8A')
    pBdr.append(bottom)
    pPr.append(pBdr)

    # Pie: texto a la izquierda, y a la derecha el prefijo + el número de
    # página real (campo de Word, no un número fijo).
    footer = sec.footer
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)
    tf = footer.add_table(rows=1, cols=2, width=Mm(170))
    tf.alignment = WD_TABLE_ALIGNMENT.CENTER
    tf.columns[0].width = Mm(90)
    tf.columns[1].width = Mm(80)
    _quitar_bordes_tabla(tf)

    r = tf.rows[0].cells[0].paragraphs[0].add_run(pie_izquierda)
    r.font.size = Pt(7.2); r.font.color.rgb = GRIS; r.font.name = 'Calibri'

    p_der = tf.rows[0].cells[1].paragraphs[0]
    p_der.alignment = AP.RIGHT
    r = p_der.add_run(pie_derecha_prefijo)
    r.font.size = Pt(7.2); r.font.color.rgb = GRIS; r.font.name = 'Calibri'
    _campo_pagina(p_der)


def _campo_pagina(paragraph):
    """Inserta el campo PAGE de Word: el número real de cada hoja, no un
    número escrito a mano que quedaría fijo si se agregan páginas."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    run = paragraph.add_run()
    run.font.size = Pt(7.2)
    run.font.color.rgb = GRIS
    campo_ini = OxmlElement('w:fldChar')
    campo_ini.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.text = 'PAGE'
    campo_fin = OxmlElement('w:fldChar')
    campo_fin.set(qn('w:fldCharType'), 'end')
    run._r.append(campo_ini)
    run._r.append(instr)
    run._r.append(campo_fin)


def escribir(bloques, destino, plantilla=None, cabecera_estudio=None):
    """Arma el .docx.

    `plantilla`: la ruta de un .docx a usar como base —el membrete.docx de la
    Prefectura, con su propia cabecera y pie ya dentro—.
    `cabecera_estudio`: si no hay plantilla, un diccionario con los argumentos
    de `agregar_cabecera_pie_estudio()` para armar la cabecera del estudio
    sobre un documento en blanco. Sin ninguno de los dos, el documento sale
    sin cabecera ni pie.
    """
    if plantilla:
        doc = Document(plantilla)
        # La plantilla trae el cuerpo vacio con un parrafo suelto: se retira
        # para que el documento no empiece con una linea en blanco.
        for p in list(doc.paragraphs):
            if not p.text.strip():
                p._element.getparent().remove(p._element)
    else:
        doc = Document()
        if cabecera_estudio:
            agregar_cabecera_pie_estudio(doc, **cabecera_estudio)
    _pintar(doc, bloques, True)
    doc.save(destino)
    return destino
