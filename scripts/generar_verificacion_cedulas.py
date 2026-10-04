# -*- coding: utf-8 -*-
"""
Solicitud a campo: confirmar la cédula de las fichas que no la tienen clara.

Salen dos archivos del mismo material, uno para trabajar y otro para leer:

* Excel — la hoja de trabajo: una fila por ficha, filtrable por comunidad, con
  las columnas de seguimiento (ESTADO, FECHA, OBSERVACIÓN).
* Word  — el pedido: qué se revisa, cómo, y cuánto hay en cada comunidad.

Qué fichas entran
-----------------
Las define `generar_listado_titulares.casos_para_revisar()`, la misma función
que alimenta la hoja PARA REVISAR del listado que se le dio al sociólogo: sin
cédula (63), con dígito verificador no válido (79) y las que comparten cédula
con otra ficha principal (500 fichas, 217 cédulas). 15 fichas tienen dos
motivos a la vez, así que el total es 627 y no 642.

Regla de diseño heredada de `generar_excel_revision_campo.py`
-------------------------------------------------------------
**El Excel NO es donde se corrige. Se corrige en QField.** Si el técnico
escribiera aquí la cédula buena, ese dato no llegaría al padrón y quedarían dos
versiones del mismo dato. Por eso las columnas de trabajo son de seguimiento y
la cédula hallada se anota en OBSERVACIÓN, como constancia.

Protección contra pérdida
-------------------------
Este Excel NO conserva las notas al regenerarse (el de revisión de campo sí, y
a cambio es mucho más complejo). En su lugar, el script se niega a sobrescribir
un archivo que ya tenga algún ESTADO distinto de «Pendiente»: sin eso, una
regeneración tras la primera semana de trabajo borraría lo revisado.

Uso:  python -X utf8 scripts/generar_verificacion_cedulas.py
      python -X utf8 scripts/generar_verificacion_cedulas.py --forzar
"""
import argparse
import collections
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

from armar_entrega_final import catalogo  # noqa: E402
from generar_listado_titulares import cargar, casos_para_revisar  # noqa: E402
from generar_fichas_pdf import (  # noqa: E402
    A4, AZUL as AZUL_PDF, TINTA, FECHA_CORTE, PROJECT_LOCATION, PROJECT_SUBTITLE,
    PROJECT_SUBTITLE_ALCANCE, PROJECT_TITLE, PUB, colors, mm, Paragraph,
    ParagraphStyle, Spacer, tabla_datos, titulo_seccion,
)
import a_word  # noqa: E402

ESCRITORIO = r'C:\Users\HP\OneDrive\Escritorio'
EXCEL = os.path.join(ESCRITORIO, 'VERIFICACION DE CEDULAS - solicitud a campo.xlsx')
WORD = os.path.join(ESCRITORIO, 'VERIFICACION DE CEDULAS - solicitud a campo.docx')

AZUL = '1E3A8A'
FILL_CAB = PatternFill('solid', fgColor=AZUL)
FILL_TRAB = PatternFill('solid', fgColor='F1F5F9')
FILL_TOT = PatternFill('solid', fgColor='EEF3EF')
LADO = Side(style='thin', color='CBD5E1')
BORDE = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)
CAB = Font(bold=True, color='FFFFFF', size=10)

ESTADOS = ['Pendiente', 'Cédula confirmada', 'Cédula corregida en QField',
           'Misma persona con varios predios', 'Ficha duplicada',
           'Persona distinta', 'No se pudo verificar (anotar por qué)']


def es_ruc_cortado(x):
    """10 dígitos con tercer dígito 6 o 9: es el comienzo del RUC de una
    organización (6 sector público, 9 jurídica privada), al que le faltan los
    tres últimos (001). No es una cédula mal digitada: es un RUC incompleto, y
    se resuelve completándolo, no corrigiéndolo."""
    return (x['estado_ced'] == 'Dígito verificador NO válido'
            and len(x['cedula']) == 10 and x['cedula'][2] in '69')


def que_revisar(x, por_numero):
    """El motivo, con el dato que más ayuda a resolverlo en el terreno: en las
    cédulas repetidas, si el resto de fichas está en la misma comunidad o en
    otra (no es lo mismo ir a un solo sitio que a varios)."""
    partes = []
    if x['estado_ced'] == 'Sin cédula':
        partes.append('Sin cédula registrada')
    elif x['estado_ced'] == 'Dígito verificador NO válido':
        partes.append('RUC de organización incompleto (10 de 13 dígitos)'
                      if es_ruc_cortado(x) else 'Cédula con dígito verificador no válido')
    if x['otras']:
        n = len(x['otras']) + 1
        coms = {int(c[5:7]) for c in x['otras']}
        otras = sorted(coms - {x['com_n']})
        if not otras:
            donde = 'todas en esta comunidad'
        elif x['com_n'] in coms:
            donde = 'en esta comunidad y en: ' + '; '.join(por_numero[c] for c in otras)
        else:
            donde = 'las otras en: ' + '; '.join(por_numero[c] for c in otras)
        partes.append(f'Misma cédula en {n} fichas principales ({donde})')
    return '. '.join(partes)


def contar_por_comunidad(casos):
    por = collections.OrderedDict()
    for x, _ in sorted(casos, key=lambda c: (c[0]['com_n'], c[0]['codigo'])):
        d = por.setdefault((x['com_n'], x['comunidad'], x['sector']),
                           {'sin': 0, 'inv': 0, 'ruc': 0, 'rep': 0, 'tot': 0})
        d['tot'] += 1
        d['sin'] += x['estado_ced'] == 'Sin cédula'
        d['ruc'] += es_ruc_cortado(x)
        d['inv'] += (x['estado_ced'] == 'Dígito verificador NO válido'
                     and not es_ruc_cortado(x))
        d['rep'] += bool(x['otras'])
    return por


def hay_trabajo_hecho(ruta):
    """¿Ya alguien marcó algo en el Excel entregado? Si sí, no se pisa."""
    if not os.path.exists(ruta):
        return False
    try:
        ws = load_workbook(ruta, read_only=True)['Verificar']
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if fila[8] not in (None, '', 'Pendiente'):
                return True
    except Exception:
        return False
    return False


def escribir_excel(casos, por, por_numero):
    wb = Workbook()
    n = len(casos)
    tot_sin = sum(d['sin'] for d in por.values())
    tot_inv = sum(d['inv'] for d in por.values())
    tot_ruc = sum(d['ruc'] for d in por.values())
    tot_rep = sum(d['rep'] for d in por.values())

    # ── Instrucciones ──
    ws = wb.active
    ws.title = 'Instrucciones'
    ws.sheet_view.showGridLines = False
    ws['A1'] = 'Verificación de cédulas — solicitud a campo'
    ws['A1'].font = Font(bold=True, size=15, color=AZUL)
    ws['A2'] = f'Padrón de usuarios · Información al {FECHA_CORTE}'
    ws['A2'].font = Font(size=10, color='64748B')
    lineas = [
        ('Qué se pide',
         f'Confirmar la cédula del titular en {n} fichas principales. Son cuatro casos y '
         'cada fila dice cuál en la columna «Qué revisar».'),
        ('1. Sin cédula',
         f'({tot_sin} fichas) Pedir la cédula al titular o verla en sus papeles.'),
        ('2. Dígito verificador no válido',
         f'({tot_inv} fichas, personas) Casi siempre es un error de digitación: revisar la '
         'ficha en papel o preguntar al titular y anotar la cédula correcta.'),
        ('3. RUC de organización incompleto',
         f'({tot_ruc} fichas) Es el RUC de una comunidad, asociación o institución escrito con '
         '10 dígitos. No es un error de digitación: hay que completarlo con los 13 (termina '
         'en 001). Pedirlo a la directiva o verlo en su documento.'),
        ('4. Misma cédula en varias fichas',
         f'({tot_rep} fichas) Averiguar cuál de estos casos es: la misma persona con varios '
         'predios, la misma ficha levantada dos veces, o dos personas distintas con la cédula '
         'mal copiada.'),
        ('Cómo se trabaja',
         'Filtre la hoja «Verificar» por su comunidad (columna B). En cada fila elija el ESTADO '
         'de la lista, anote la FECHA y escriba en OBSERVACIÓN lo que encontró; si halló la '
         'cédula correcta, escríbala ahí.'),
        ('Qué NO se hace aquí',
         'El Excel no es donde se corrige. La cédula se corrige en QField, una vez confirmada. '
         'Estas columnas dejan constancia de lo revisado.'),
        ('Fichas con dos motivos',
         f'{tot_sin + tot_inv + tot_ruc + tot_rep - n} fichas tienen dos motivos a la vez; por eso el total '
         f'de filas ({n}) es menor que la suma de los cuatro casos.'),
    ]
    r = 4
    for t, txt in lineas:
        ws.cell(row=r, column=1, value=t).font = Font(bold=True, size=11, color=AZUL)
        c = ws.cell(row=r, column=2, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.font = Font(size=10.5)
        ws.row_dimensions[r].height = 46
        r += 1
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 95

    # ── Resumen por comunidad ──
    ws = wb.create_sheet('Resumen')
    cab = ['N°', 'Comunidad', 'Sector', 'Sin cédula', 'Dígito no válido',
           'RUC incompleto', 'Cédula repetida', 'Fichas a verificar']
    ws.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for (num, nom, sec), d in por.items():
        ws.append([num, nom, sec, d['sin'], d['inv'], d['ruc'], d['rep'], d['tot']])
    ws.append(['', 'TOTAL', '', tot_sin, tot_inv, tot_ruc, tot_rep, n])
    for rr in range(2, ws.max_row + 1):
        for cc in range(1, len(cab) + 1):
            c = ws.cell(row=rr, column=cc)
            c.border = BORDE
            c.font = Font(size=10)
    for cc in range(1, len(cab) + 1):
        c = ws.cell(row=ws.max_row, column=cc)
        c.font, c.fill = Font(bold=True, size=10), FILL_TOT
    for i, w in enumerate([6, 34, 10, 12, 14, 14, 14, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30

    # ── Verificar ──
    ws = wb.create_sheet('Verificar')
    # ESTADO queda en la columna 9 a propósito: hay_trabajo_hecho() la lee por posición.
    cab = ['Código', 'Comunidad', 'Clave catastral', 'Titular', 'Cédula registrada',
           'Qué revisar', 'Levantó', 'Sector', 'ESTADO', 'Fecha revisión', 'OBSERVACIÓN']
    ws.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    orden = sorted(casos, key=lambda c: (c[0]['com_n'], c[0]['cedula'], c[0]['codigo']))
    for x, _ in orden:
        ws.append([x['codigo'], x['comunidad'], x['clave'], x['nombre'], x['cedula'],
                   que_revisar(x, por_numero), x['tecnico'], x['sector'],
                   'Pendiente', None, None])
    for rr in range(2, ws.max_row + 1):
        for cc in range(1, len(cab) + 1):
            c = ws.cell(row=rr, column=cc)
            c.border = BORDE
            c.font = Font(size=10)
            c.alignment = Alignment(vertical='top', wrap_text=(cc in (4, 6, 11)))
            if cc in (3, 5):
                c.number_format = '@'      # clave y cédula como texto
            if cc in (9, 10, 11):
                c.fill = FILL_TRAB
    dv = DataValidation(type='list', formula1='"{}"'.format(','.join(ESTADOS)),
                        allow_blank=True, showDropDown=False)
    dv.error = 'Elige una opción de la lista'
    ws.add_data_validation(dv)
    dv.add(f'I2:I{ws.max_row}')
    for i, w in enumerate([18, 30, 17, 36, 15, 58, 16, 10, 30, 14, 48], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f'A1:K{ws.max_row}'
    ws.freeze_panes = 'D2'
    ws.row_dimensions[1].height = 30

    wb.save(EXCEL)

    # Se relee lo guardado y se cuadra contra la lista de origen.
    chk = load_workbook(EXCEL, read_only=True)
    filas = chk['Verificar'].max_row - 1
    suma = sum(f[7] for f in chk['Resumen'].iter_rows(min_row=2, values_only=True)
               if f[1] != 'TOTAL')
    assert filas == n, (filas, n)
    assert suma == n, (suma, n)


def escribir_word(casos, por):
    n = len(casos)
    tot_sin = sum(d['sin'] for d in por.values())
    tot_inv = sum(d['inv'] for d in por.values())
    tot_ruc = sum(d['ruc'] for d in por.values())
    tot_rep = sum(d['rep'] for d in por.values())
    en_rep = len({x['cedula'] for x, _ in casos if x['otras']})

    st_h = ParagraphStyle('h', fontName='Helvetica-Bold', fontSize=13, leading=16,
                          textColor=AZUL_PDF, alignment=1)
    st_sub = ParagraphStyle('sub', fontName='Helvetica', fontSize=9.5, leading=12,
                            alignment=1, textColor=colors.HexColor('#475569'))
    st_p = ParagraphStyle('p', fontName='Helvetica', fontSize=9.5, leading=13,
                          textColor=TINTA, alignment=4)
    st_li = ParagraphStyle('li', parent=st_p, leftIndent=6 * mm, spaceAfter=2)

    b = [Paragraph('Verificación de cédulas', st_h), Spacer(0, 1.5 * mm),
         Paragraph('Solicitud a campo · Padrón de usuarios', st_sub), Spacer(0, 5 * mm),
         Paragraph(f'Se pide confirmar la cédula del titular en <b>{n} fichas principales</b> '
                   'del padrón. El objetivo es dejar cada ficha con una cédula verificada. '
                   'Las fichas están en el Excel adjunto, ordenadas por comunidad, con la '
                   'columna «Qué revisar» que dice el caso de cada una.', st_p)]

    b += titulo_seccion('1. Qué hay que revisar')
    b.append(tabla_datos(
        ['Caso', 'Fichas', 'Qué se pide'],
        [['Sin cédula registrada', str(tot_sin),
          'Pedir la cédula al titular o verla en sus papeles.'],
         ['Dígito verificador no válido (personas)', str(tot_inv),
          'Revisar la ficha en papel o preguntar al titular: casi siempre es un error de '
          'digitación.'],
         ['RUC de organización incompleto', str(tot_ruc),
          'Comunidades, asociaciones o instituciones cuyo RUC está escrito con 10 dígitos. '
          'Completarlo con los 13 (termina en 001): pedirlo a la directiva o verlo en su '
          'documento.'],
         ['Misma cédula en varias fichas principales', str(tot_rep),
          f'Son {en_rep} cédulas. Averiguar si es la misma persona con varios predios, la '
          'misma ficha levantada dos veces, o dos personas con la cédula mal copiada.']],
        [58 * mm, 18 * mm, 94 * mm]))
    b.append(Spacer(0, 1.5 * mm))
    b.append(Paragraph(f'{tot_sin + tot_inv + tot_ruc + tot_rep - n} fichas tienen dos motivos a la vez, '
                       f'por eso el total es {n} y no la suma de los cuatro casos.', st_p))

    b += titulo_seccion('2. Cómo se trabaja')
    for t in [
        '<b>1.</b> Abrir el Excel y filtrar la hoja «Verificar» por la comunidad que toca '
        '(columna «Comunidad»).',
        '<b>2.</b> En cada fila, confirmar la cédula con el titular o con su documento.',
        '<b>3.</b> Elegir el ESTADO de la lista, poner la fecha y anotar en OBSERVACIÓN lo '
        'que se encontró; si se halló la cédula correcta, escribirla ahí.',
        '<b>4.</b> Corregir la cédula en QField una vez confirmada.',
    ]:
        b.append(Paragraph(t, st_li))
    b.append(Spacer(0, 2 * mm))
    b.append(Paragraph('El Excel no es el lugar donde se corrige el dato: deja constancia de '
                       'lo revisado. La corrección se hace en QField, de modo que el padrón '
                       'tenga una sola versión de cada cédula.', st_p))

    b += titulo_seccion('3. Cuántas fichas hay en cada comunidad')
    filas = [[nom, str(d['sin']), str(d['inv']), str(d['ruc']), str(d['rep']), str(d['tot'])]
             for (_, nom, _), d in por.items()]
    filas.append(['TOTAL', str(tot_sin), str(tot_inv), str(tot_ruc), str(tot_rep), str(n)])
    b.append(tabla_datos(['Comunidad', 'Sin cédula', 'Dígito no válido', 'RUC incompleto',
                          'Cédula repetida', 'Fichas a verificar'],
                         filas, [52 * mm, 22 * mm, 24 * mm, 24 * mm, 24 * mm, 24 * mm],
                         ultima_negrita=True))

    b += titulo_seccion('4. Qué se devuelve')
    b.append(Paragraph('El mismo Excel, con ESTADO, FECHA y OBSERVACIÓN llenos en las fichas '
                       'revisadas. Las que no se alcancen a verificar se dejan en «Pendiente» '
                       'o con «No se pudo verificar» y el motivo.', st_p))

    a_word.escribir(b, WORD, cabecera_estudio=dict(
        titulo=PROJECT_TITLE, subtitulo=PROJECT_SUBTITLE,
        alcance=PROJECT_SUBTITLE_ALCANCE, ubicacion=PROJECT_LOCATION,
        logo_izq=os.path.join(PUB, 'logo-izq.png'),
        logo_der=os.path.join(PUB, 'logo-der.png'),
        pie_izquierda='AP&CATASTROS',
        pie_derecha_prefijo=f'CONSORCIO CAYAMBE SPT · Datos al {FECHA_CORTE} · pág. '))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--forzar', action='store_true',
                    help='sobrescribir aunque el Excel ya tenga revisiones marcadas')
    args = ap.parse_args()

    if hay_trabajo_hecho(EXCEL) and not args.forzar:
        print('El Excel ya tiene fichas revisadas (ESTADO distinto de «Pendiente»).')
        print('No se sobrescribe: se perdería ese trabajo. Use --forzar solo si es a propósito.')
        return 1

    filas, _ = cargar()
    _, por_numero = catalogo()
    casos = casos_para_revisar(filas)
    por = contar_por_comunidad(casos)

    escribir_excel(casos, por, por_numero)
    escribir_word(casos, por)

    print(f'Fichas a verificar : {len(casos)}')
    print(f'  sin cédula       : {sum(d["sin"] for d in por.values())}')
    print(f'  dígito no válido : {sum(d["inv"] for d in por.values())}')
    print(f'  RUC incompleto   : {sum(d["ruc"] for d in por.values())}')
    print(f'  cédula repetida  : {sum(d["rep"] for d in por.values())}')
    print(f'  comunidades      : {len(por)}')
    print('✔', EXCEL)
    print('✔', WORD)
    return 0


if __name__ == '__main__':
    sys.exit(main())
