# -*- coding: utf-8 -*-
"""
Segunda solicitud a campo de cédulas: solo lo que la oficina no pudo resolver.

La solicitud del 4-oct (`generar_verificacion_cedulas.py`) pedía 627 fichas.
El 7-oct se revisaron todas en oficina (catastro municipal, otras fichas del
mismo titular, dígito verificador, RUC completado con 001) y 511 quedaron
resueltas. Este script arma el pedido de las que quedan, que REEMPLAZA al del
4-oct: el equipo de campo no debe visitar las 511.

De dónde sale la lista
----------------------
Las fichas pendientes son las filas de la hoja «Solicitud 4-oct» de
`VERIFICACIONES EN OFICINA (7-oct-2026).xlsx` cuyo «Resultado» dice «requiere
campo». El motivo, la clave y el técnico se vuelven a leer de la base con
`casos_para_revisar()`, la misma función de la solicitud del 4-oct: si una de
esas fichas ya no es caso (alguien corrigió la cédula en QField), el script se
detiene y lo dice, en vez de pedir una visita que sobra.

Además lleva una hoja «Llamadas»: cuatro casos fuera de la lista de cédulas que
se confirman por teléfono (la hoja homónima del archivo de oficina).

Misma regla que la solicitud del 4-oct: el Excel no es donde se corrige; la
cédula se corrige en QField y aquí solo se deja constancia. Y la misma
protección: no se sobrescribe un Excel que ya tenga algún ESTADO marcado.

Uso:  python -X utf8 scripts/generar_verificacion_cedulas_campo2.py [--forzar]
"""
import argparse
import collections
import os
import re
import sys
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

from armar_entrega_final import catalogo  # noqa: E402
from excel_compat import aplicar_formatos  # noqa: E402
from generar_listado_titulares import cargar, casos_para_revisar  # noqa: E402
from generar_fichas_pdf import (  # noqa: E402
    AZUL as AZUL_PDF, TINTA, FECHA_CORTE, PROJECT_LOCATION, PROJECT_SUBTITLE,
    PROJECT_SUBTITLE_ALCANCE, PROJECT_TITLE, PUB, colors, mm, Paragraph,
    ParagraphStyle, Spacer, tabla_datos, titulo_seccion,
)
from generar_verificacion_cedulas import (  # noqa: E402
    AZUL, BORDE, CAB, ESTADOS, FILL_CAB, FILL_TOT, FILL_TRAB,
    contar_por_comunidad, es_ruc_cortado, hay_trabajo_hecho, que_revisar,
)
from informe_estilo import esn  # noqa: E402
import a_word  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(AQUI))
OFICINA = os.path.join(RAIZ, 'CONTRATO Y FISCALIZACION', 'RESPUESTA OBSERVACIONES PRODUCTO 5',
                       'VERIFICACIONES EN OFICINA (7-oct-2026).xlsx')
ESCRITORIO = r'C:\Users\HP\OneDrive\Escritorio'
NOMBRE = 'VERIFICACION DE CEDULAS - 116 fichas a campo (11-oct-2026)'
EXCEL = os.path.join(ESCRITORIO, NOMBRE + '.xlsx')
WORD = os.path.join(ESCRITORIO, NOMBRE + '.docx')

FECHA_SOLICITUD = '11 de octubre de 2026'
ESTADOS_LLAMADA = ['Pendiente', 'Confirmado por teléfono', 'No contesta',
                   'Requiere visita (anotar por qué)']


def leer_oficina():
    """Las fichas que siguen pendientes, las resueltas y las llamadas."""
    wb = load_workbook(OFICINA, read_only=True)
    pend, resueltas = collections.OrderedDict(), set()
    for f in wb['Solicitud 4-oct'].iter_rows(min_row=5, values_only=True):
        if not f[0]:
            continue
        if 'requiere campo' in str(f[5] or '').lower():
            pend[f[0]] = {'resultado': f[5], 'telefono': f[8]}
        else:
            resueltas.add(f[0])
    llamadas = [f for f in wb['Llamadas'].iter_rows(min_row=5, values_only=True) if f[0]]
    return pend, resueltas, llamadas


def en_castellano(txt):
    """Las áreas de la hoja de llamadas se escribieron en formato anglosajón
    («10941.0 m²») y las fechas como 2026-05-28; aquí pasan a castellano."""
    txt = re.sub(r'(\d+(?:\.\d+)?) m²', lambda m: esn(float(m.group(1)), 2) + ' m²', txt)
    return re.sub(r'(\d{4})-(\d{2})-(\d{2})', r'\3/\2/\1', txt)


def que_revisar_campo(x, por_numero, pend, resueltas):
    """El motivo de la solicitud del 4-oct, más dos datos que ayudan en el
    terreno: qué otras fichas comparten la cédula (y si ya se resolvieron) y,
    cuando la oficina halló más de una cédula posible, que es eso lo que hay
    que decidir."""
    t = que_revisar(x, por_numero)
    if x['otras']:
        t += '. Otras fichas con esa cédula: ' + ', '.join(
            c + (' (resuelta en oficina)' if c in resueltas else '')
            for c in sorted(x['otras']))
    if 'Varias' in str(pend[x['codigo']]['resultado']):
        t += '. En oficina aparecieron varias cédulas posibles: confirmar cuál es'
    return t


def _nombre(x):
    """El nombre sin dobles espacios ni tildes, para agrupar las variantes con
    que se escribió el mismo titular («PITANA» / «PITANÁ»)."""
    s = unicodedata.normalize('NFKD', ' '.join(x['nombre'].upper().split()))
    return ''.join(ch for ch in s if not unicodedata.combining(ch))


def titulares(casos):
    """Personas u organizaciones distintas a contactar. El código S-C-R agrupa
    por cédula, y aquí hay cédulas que comparten dos personas (justo lo que se
    va a verificar): por eso se cuenta código R más nombre."""
    return len({(x['codigo'][:12], _nombre(x)) for x, _ in casos})


def organizaciones_alpaka(casos):
    """Los titulares de Alpaka con varias fichas: son organizaciones (comités,
    comunidades, una empresa) sin cédula ni RUC completo, y una consulta a su
    directiva resuelve todas sus fichas de una vez."""
    grupos = collections.defaultdict(list)
    for x, _ in casos:
        if 'ALPAKA' in x['comunidad']:
            grupos[x['codigo'][:12]].append(x)
    salida = []
    for g in grupos.values():
        if len(g) < 2:
            continue
        nombre = collections.Counter(' '.join(x['nombre'].split()) for x in g).most_common(1)[0][0]
        if any(es_ruc_cortado(x) for x in g):
            falta = 'RUC incompleto (10 de 13 dígitos)'
        elif all(x['estado_ced'] == 'Sin cédula' for x in g):
            falta = 'Sin cédula ni RUC'
        else:
            falta = 'Cédula a confirmar'
        salida.append((nombre, len(g), falta))
    return sorted(salida, key=lambda t: -t[1])


def escribir_excel(casos, por, por_numero, pend, resueltas, llamadas):
    wb = Workbook()
    n = len(casos)
    tot = {k: sum(d[k] for d in por.values()) for k in ('sin', 'inv', 'ruc', 'rep')}
    n_tit = titulares(casos)
    orgs = organizaciones_alpaka(casos)
    n_alp = por_comunidad_n(por, 'ALPAKA')
    n_tit_alp = titulares([c for c in casos if 'ALPAKA' in c[0]['comunidad']])

    # ── Instrucciones ──
    ws = wb.active
    ws.title = 'Instrucciones'
    ws.sheet_view.showGridLines = False
    ws['A1'] = f'Verificación de cédulas — {n} fichas a campo'
    ws['A1'].font = Font(bold=True, size=15, color=AZUL)
    ws['A2'] = (f'Padrón de usuarios · Solicitud del {FECHA_SOLICITUD} · '
                f'Información al {FECHA_CORTE}')
    ws['A2'].font = Font(size=10, color='64748B')
    lineas = [
        ('REEMPLAZA a la del 4-oct',
         f'La solicitud del 4 de octubre pedía 627 fichas. En oficina se resolvieron '
         f'{len(resueltas)} con el catastro municipal, otras fichas del mismo titular y el dígito '
         f'verificador: esas NO se visitan. Quedan estas {n}. Use este archivo y deje de usar '
         'el del 4 de octubre.'),
        ('Qué se pide',
         f'Confirmar la cédula del titular en {n} fichas principales, que corresponden a {n_tit} '
         'titulares (personas u organizaciones). La columna «Qué revisar» dice el caso de cada '
         'ficha.'),
        ('1. Sin cédula',
         f'({tot["sin"]} fichas) Pedir la cédula al titular o verla en sus papeles.'),
        ('2. Dígito verificador no válido',
         f'({tot["inv"]} fichas, personas) Casi siempre es un error de digitación: revisar la '
         'ficha en papel o preguntar al titular y anotar la cédula correcta.'),
        ('3. RUC de organización incompleto',
         f'({tot["ruc"]} fichas) Es el RUC de una comunidad, asociación o institución escrito con '
         '10 dígitos. Hay que completarlo con los 13 (termina en 001): pedirlo a la directiva '
         'o verlo en su documento.'),
        ('4. Misma cédula en varias fichas',
         f'({tot["rep"]} fichas) Averiguar si es la misma persona con varios predios, la misma '
         'ficha levantada dos veces, o dos personas distintas con la cédula mal copiada. La '
         'columna «Qué revisar» lista las otras fichas con esa cédula.'),
        ('Alpaka',
         f'Tiene {n_alp} de las {n} fichas, pero son solo {n_tit_alp} titulares. '
         f'{sum(o[1] for o in orgs)} fichas son de {len(orgs)} organizaciones: '
         + '; '.join(f'{o[0]} ({o[1]})' for o in orgs)
         + '. Con una consulta a cada directiva se resuelven todas sus fichas.'),
        ('Cómo se trabaja',
         'Filtre la hoja «Verificar» por su comunidad (columna B). En cada fila elija el ESTADO '
         'de la lista, anote la FECHA y escriba en OBSERVACIÓN lo que encontró; si halló la '
         'cédula correcta, escríbala ahí. El teléfono es el que dio el titular en la entrevista.'),
        ('Llamadas',
         f'La hoja «Llamadas» tiene {len(llamadas)} casos que no necesitan visita: se confirman '
         'por teléfono.'),
        ('Qué NO se hace aquí',
         'El Excel no es donde se corrige. La cédula se corrige en QField, una vez confirmada. '
         'Estas columnas dejan constancia de lo revisado.'),
        ('Uso del archivo',
         'Trae teléfonos y cédulas: es de uso interno del equipo y no se comparte fuera de él.'),
    ]
    r = 4
    for t, txt in lineas:
        ws.cell(row=r, column=1, value=t).font = Font(bold=True, size=11, color=AZUL)
        c = ws.cell(row=r, column=2, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.font = Font(size=10.5)
        ws.cell(row=r, column=1).alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[r].height = 46
        r += 1
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 95

    # ── Resumen por comunidad ──
    ws = wb.create_sheet('Resumen')
    cab = ['N°', 'Comunidad', 'Sector', 'Sin cédula', 'Dígito no válido',
           'RUC incompleto', 'Cédula repetida', 'Fichas a verificar', 'Titulares']
    ws.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    tit_com = collections.defaultdict(set)
    for x, _ in casos:
        tit_com[x['com_n']].add(x['codigo'][:12])
    for (num, nom, sec), d in por.items():
        ws.append([num, nom, sec, d['sin'], d['inv'], d['ruc'], d['rep'], d['tot'],
                   len(tit_com[num])])
    ws.append(['', 'TOTAL', '', tot['sin'], tot['inv'], tot['ruc'], tot['rep'], n, n_tit])
    for rr in range(2, ws.max_row + 1):
        for cc in range(1, len(cab) + 1):
            c = ws.cell(row=rr, column=cc)
            c.border = BORDE
            c.font = Font(size=10)
    for cc in range(1, len(cab) + 1):
        c = ws.cell(row=ws.max_row, column=cc)
        c.font, c.fill = Font(bold=True, size=10), FILL_TOT
    for i, w in enumerate([6, 34, 10, 12, 14, 14, 14, 16, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30

    # ── Verificar ──
    ws = wb.create_sheet('Verificar')
    # ESTADO en la columna 9, como en la solicitud del 4-oct: hay_trabajo_hecho() la lee ahí.
    cab = ['Código', 'Comunidad', 'Clave catastral', 'Titular', 'Cédula registrada',
           'Qué revisar', 'Teléfono', 'Levantó', 'ESTADO', 'Fecha revisión', 'OBSERVACIÓN']
    ws.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    orden = sorted(casos, key=lambda c: (c[0]['com_n'], c[0]['codigo']))
    for x, _ in orden:
        ws.append([x['codigo'], x['comunidad'], x['clave'], x['nombre'], x['cedula'],
                   que_revisar_campo(x, por_numero, pend, resueltas),
                   pend[x['codigo']]['telefono'], x['tecnico'], 'Pendiente', None, None])
    for rr in range(2, ws.max_row + 1):
        for cc in range(1, len(cab) + 1):
            c = ws.cell(row=rr, column=cc)
            c.border = BORDE
            c.font = Font(size=10)
            c.alignment = Alignment(vertical='top', wrap_text=(cc in (4, 6, 11)))
            if cc in (3, 5, 7):
                c.number_format = '@'      # clave, cédula y teléfono como texto
            if cc in (9, 10, 11):
                c.fill = FILL_TRAB
    dv = DataValidation(type='list', formula1='"{}"'.format(','.join(ESTADOS)),
                        allow_blank=True, showDropDown=False)
    dv.error = 'Elige una opción de la lista'
    ws.add_data_validation(dv)
    dv.add(f'I2:I{ws.max_row}')
    for i, w in enumerate([18, 30, 17, 36, 15, 62, 13, 16, 30, 14, 48], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f'A1:K{ws.max_row}'
    ws.freeze_panes = 'D2'
    ws.row_dimensions[1].height = 30

    # ── Llamadas ──
    ws = wb.create_sheet('Llamadas')
    cab = ['Código', 'Titular', 'Teléfono', 'Comunidad', 'Qué preguntar',
           'ESTADO', 'Fecha llamada', 'RESPUESTA']
    ws.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for f in llamadas:
        ws.append([f[0], f[1], f[2], f[3], en_castellano(f[4]), 'Pendiente', None, None])
    for rr in range(2, ws.max_row + 1):
        ws.row_dimensions[rr].height = 48
        for cc in range(1, len(cab) + 1):
            c = ws.cell(row=rr, column=cc)
            c.border = BORDE
            c.font = Font(size=10)
            c.alignment = Alignment(vertical='top', wrap_text=(cc in (2, 5, 8)))
            if cc == 3:
                c.number_format = '@'
            if cc in (6, 7, 8):
                c.fill = FILL_TRAB
    dv = DataValidation(type='list', formula1='"{}"'.format(','.join(ESTADOS_LLAMADA)),
                        allow_blank=True, showDropDown=False)
    dv.error = 'Elige una opción de la lista'
    ws.add_data_validation(dv)
    dv.add(f'F2:F{ws.max_row}')
    for i, w in enumerate([18, 30, 13, 16, 70, 26, 14, 48], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C2'
    ws.row_dimensions[1].height = 30

    wb.save(EXCEL)
    aplicar_formatos(EXCEL)     # sin esto el Excel local no pinta colores ni formatos

    chk = load_workbook(EXCEL, read_only=True)
    filas = chk['Verificar'].max_row - 1
    suma = sum(f[7] for f in chk['Resumen'].iter_rows(min_row=2, values_only=True)
               if f[1] != 'TOTAL')
    assert filas == n, (filas, n)
    assert suma == n, (suma, n)


def por_comunidad_n(por, nombre):
    return sum(d['tot'] for (_, nom, _), d in por.items() if nombre in nom)


def escribir_word(casos, por, resueltas, llamadas):
    n = len(casos)
    tot = {k: sum(d[k] for d in por.values()) for k in ('sin', 'inv', 'ruc', 'rep')}
    n_tit = titulares(casos)
    en_rep = len({x['cedula'] for x, _ in casos if x['otras']})

    st_h = ParagraphStyle('h', fontName='Helvetica-Bold', fontSize=13, leading=16,
                          textColor=AZUL_PDF, alignment=1)
    st_sub = ParagraphStyle('sub', fontName='Helvetica', fontSize=9.5, leading=12,
                            alignment=1, textColor=colors.HexColor('#475569'))
    st_p = ParagraphStyle('p', fontName='Helvetica', fontSize=9.5, leading=13,
                          textColor=TINTA, alignment=4)
    st_li = ParagraphStyle('li', parent=st_p, leftIndent=6 * mm, spaceAfter=2)

    b = [Paragraph('Verificación de cédulas', st_h), Spacer(0, 1.5 * mm),
         Paragraph(f'Segunda solicitud a campo · {n} fichas · {FECHA_SOLICITUD}', st_sub),
         Spacer(0, 5 * mm),
         Paragraph(f'<b>Esta solicitud reemplaza a la del 4 de octubre.</b> Aquella pedía '
                   f'confirmar la cédula en 627 fichas principales. En oficina se resolvieron '
                   f'{esn(len(resueltas))} con el catastro municipal, las otras fichas del mismo '
                   'titular y el dígito verificador, y no hace falta visitarlas. Quedan '
                   f'<b>{n} fichas</b>, que corresponden a <b>{n_tit} titulares</b> (personas u '
                   'organizaciones). Están en el Excel adjunto, ordenadas por comunidad, con la '
                   'columna «Qué revisar» que dice el caso de cada una.', st_p)]

    b += titulo_seccion('1. Qué hay que revisar')
    b.append(tabla_datos(
        ['Caso', 'Fichas', 'Qué se pide'],
        [['Sin cédula registrada', str(tot['sin']),
          'Pedir la cédula al titular o verla en sus papeles.'],
         ['Dígito verificador no válido (personas)', str(tot['inv']),
          'Revisar la ficha en papel o preguntar al titular: casi siempre es un error de '
          'digitación.'],
         ['RUC de organización incompleto', str(tot['ruc']),
          'Comunidades, asociaciones o instituciones cuyo RUC está escrito con 10 dígitos. '
          'Completarlo con los 13 (termina en 001): pedirlo a la directiva o verlo en su '
          'documento.'],
         ['Misma cédula en varias fichas principales', str(tot['rep']),
          f'Son {en_rep} cédulas. Averiguar si es la misma persona con varios predios, la '
          'misma ficha levantada dos veces, o dos personas con la cédula mal copiada.']],
        [58 * mm, 18 * mm, 94 * mm]))
    b.append(Spacer(0, 1.5 * mm))
    dobles = sum(tot.values()) - n
    if dobles:
        b.append(Paragraph(f'{dobles} fichas tienen dos motivos a la vez, por eso el total es '
                           f'{n} y no la suma de los cuatro casos.', st_p))

    b += titulo_seccion('2. Dónde está el trabajo')
    alp = por_comunidad_n(por, 'ALPAKA')
    orgs = organizaciones_alpaka(casos)
    n_org = sum(o[1] for o in orgs)
    n_tit_alp = titulares([c for c in casos if 'ALPAKA' in c[0]['comunidad']])
    b.append(Paragraph(
        f'Alpaka concentra {alp} de las {n} fichas, pero son solo {n_tit_alp} titulares. '
        f'{n_org} de esas fichas pertenecen a {len(orgs)} organizaciones que no tienen la '
        'cédula o el RUC completo: con una consulta a cada directiva se resuelven todas sus '
        f'fichas. El resto ({n - alp} fichas) se reparte en {len(por) - 1} comunidades, con '
        'pocas fichas en cada una.', st_p))
    b.append(Spacer(0, 2 * mm))
    b.append(tabla_datos(['Organización (Alpaka)', 'Fichas', 'Qué falta'],
                         [[o[0], str(o[1]), o[2]] for o in orgs],
                         [92 * mm, 18 * mm, 60 * mm]))
    b.append(Spacer(0, 3 * mm))
    b.append(Spacer(0, 2 * mm))
    filas = [[nom, str(d['sin']), str(d['inv']), str(d['ruc']), str(d['rep']), str(d['tot'])]
             for (_, nom, _), d in por.items()]
    filas.append(['TOTAL', str(tot['sin']), str(tot['inv']), str(tot['ruc']), str(tot['rep']),
                  str(n)])
    b.append(tabla_datos(['Comunidad', 'Sin cédula', 'Dígito no válido', 'RUC incompleto',
                          'Cédula repetida', 'Fichas a verificar'],
                         filas, [52 * mm, 22 * mm, 24 * mm, 24 * mm, 24 * mm, 24 * mm],
                         ultima_negrita=True))

    b += titulo_seccion('3. Cómo se trabaja')
    for t in [
        '<b>1.</b> Abrir el Excel y filtrar la hoja «Verificar» por la comunidad que toca '
        '(columna «Comunidad»).',
        '<b>2.</b> En cada fila, confirmar la cédula con el titular o con su documento. El '
        'Excel trae el teléfono que el titular dio en la entrevista.',
        '<b>3.</b> Elegir el ESTADO de la lista, poner la fecha y anotar en OBSERVACIÓN lo '
        'que se encontró; si se halló la cédula correcta, escribirla ahí.',
        '<b>4.</b> Corregir la cédula en QField una vez confirmada.',
    ]:
        b.append(Paragraph(t, st_li))
    b.append(Spacer(0, 2 * mm))
    b.append(Paragraph('El Excel no es el lugar donde se corrige el dato: deja constancia de '
                       'lo revisado. La corrección se hace en QField, de modo que el padrón '
                       'tenga una sola versión de cada cédula.', st_p))

    b += titulo_seccion('4. Casos que se resuelven por teléfono')
    b.append(Paragraph(
        f'Fuera de la lista de cédulas hay {len(llamadas)} casos que no necesitan visita. '
        'Están en la hoja «Llamadas» del Excel, con el teléfono y lo que hay que preguntar.',
        st_p))
    b.append(Spacer(0, 2 * mm))
    b.append(tabla_datos(
        ['Ficha', 'Titular', 'Qué preguntar'],
        [[f[0], f[1], en_castellano(f[4])] for f in llamadas],
        [32 * mm, 48 * mm, 90 * mm]))

    b += titulo_seccion('5. Qué se devuelve')
    b.append(Paragraph('El mismo Excel, con ESTADO, FECHA y OBSERVACIÓN llenos en las fichas '
                       'revisadas, y la hoja «Llamadas» con la respuesta de cada una. Las que '
                       'no se alcancen a verificar se dejan en «Pendiente» o con «No se pudo '
                       'verificar» y el motivo.', st_p))

    a_word.escribir(b, WORD, cabecera_estudio=dict(
        titulo=PROJECT_TITLE, subtitulo=PROJECT_SUBTITLE,
        alcance=PROJECT_SUBTITLE_ALCANCE, ubicacion=PROJECT_LOCATION,
        logo_izq=os.path.join(PUB, 'logo-izq.png'),
        logo_der=os.path.join(PUB, 'logo-der.png'),
        pie_izquierda='AP&CATASTROS',
        pie_derecha_prefijo=f'CONSORCIO CAYAMBE SPT · Datos al {FECHA_CORTE} · pág. '))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--forzar', action='store_true',
                    help='sobrescribir aunque el Excel ya tenga revisiones marcadas')
    args = ap.parse_args()

    if hay_trabajo_hecho(EXCEL) and not args.forzar:
        print('El Excel ya tiene fichas revisadas (ESTADO distinto de «Pendiente»).')
        print('No se sobrescribe: se perdería ese trabajo. Use --forzar solo si es a propósito.')
        return 1

    pend, resueltas, llamadas = leer_oficina()
    filas, _ = cargar()
    _, por_numero = catalogo()
    idx = {x['codigo']: (x, m) for x, m in casos_para_revisar(filas)}
    ya_no = [c for c in pend if c not in idx]
    if ya_no:
        print('Estas fichas ya no son caso en la base (¿se corrigieron en QField?):')
        for c in ya_no:
            print('  ', c)
        print('Actualice VERIFICACIONES EN OFICINA antes de generar la solicitud.')
        return 1
    casos = [idx[c] for c in pend]
    por = contar_por_comunidad(casos)

    escribir_excel(casos, por, por_numero, pend, resueltas, llamadas)
    escribir_word(casos, por, resueltas, llamadas)

    print(f'Fichas a verificar : {len(casos)}  ({titulares(casos)} titulares)')
    print(f'  resueltas en oficina (no se visitan): {len(resueltas)}')
    for k, t in (('sin', 'sin cédula'), ('inv', 'dígito no válido'),
                 ('ruc', 'RUC incompleto'), ('rep', 'cédula repetida')):
        print(f'  {t:<17}: {sum(d[k] for d in por.values())}')
    print(f'  comunidades      : {len(por)}')
    print(f'  llamadas         : {len(llamadas)}')
    print('✔', EXCEL)
    print('✔', WORD)
    return 0


if __name__ == '__main__':
    sys.exit(main())
