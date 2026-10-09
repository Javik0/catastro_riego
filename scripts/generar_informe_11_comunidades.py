# -*- coding: utf-8 -*-
"""
Informe de las 11 comunidades del sistema Porotog que pidió Armando (8-oct-2026):
superficie con riego y sin riego, y área sembrada en cinco clases de cultivo.
Sale en Excel (las tablas) y en Word con el membrete de la Prefectura; el PDF
se exporta del Word con el Word de la máquina.

Las 11 comunidades
------------------
Son las del informe del 18-sep (`generar_informe_comunidades_armando.py`) sin
la hacienda Guanguilquí y sin repetir COMUNA POROTOG, que allí salía dos veces
(sola como «San Vicente de Porotog» y dentro de «Porotog»). Su área sembrada
suma lo mismo que el informe agrícola del Producto 6 cita para «once
comunidades» (observación C13 del informe JAFG-2026-002).

De dónde sale cada dato
-----------------------
* Fichas: `fichas_predios.geojson` (principal / adicional).
* Superficie catastral, con riego y sin riego: `superficie_por_comunidad.json`
  (fuente única; cada predio medido una vez por su polígono, regla 12).
* Área sembrada: `cultivos.json`, superficie declarada en la sección de
  cultivos de cada ficha. Es DECLARADA: no se compara en porcentaje con la
  catastral, porque en los terrenos familiares varios herederos declaran el
  mismo terreno.
* Clases de cultivo (pedido de Armando): Pasto (mejorado y no mejorado),
  Cebolla, Flores, Papa y Otros. «Otros» incluye bosque, monte y baldío para
  que el total sea el área declarada completa; su detalle va en una hoja aparte.

Uso:  python -X utf8 scripts/generar_informe_11_comunidades.py
"""
import collections
import io
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
from excel_compat import aplicar_formatos  # noqa: E402
from informe_estilo import FECHA_CORTE, esn  # noqa: E402

BASE = os.path.abspath(os.path.join(AQUI, '..'))
GEO = os.path.join(BASE, 'public', 'geo')
ESCRITORIO = r'C:\Users\HP\OneDrive\Escritorio'
MEMBRETE = os.path.join(ESCRITORIO, 'CAYAMBE CATASTRO RIEGO', 'membrete.docx')
NOMBRE = 'INFORME 11 COMUNIDADES - riego y cultivos'
XLSX = os.path.join(ESCRITORIO, NOMBRE + '.xlsx')
DOCX = os.path.join(ESCRITORIO, NOMBRE + '.docx')

# (etiqueta, comunidad tal como está en los datos, nota)
COMUNIDADES = [
    ('Larcachaca', 'LARCACHACA', None),
    ('La Libertad', 'LA LIBERTAD', None),
    ('San Antonio', 'SAN ANTONIO', None),
    ('San José', 'SAN JOSÉ', None),
    ('Milagro', 'MILAGRO', None),
    ('La Candelaria', 'LA CANDELARIA', None),
    ('Chambitola', 'CHAMBITOLA', None),
    ('17 de Junio', 'ASOCIACIÓN 17 DE JUNIO', None),
    ('Comuna Porotog', 'COMUNA POROTOG', 'Es la que se nombró «San Vicente de Porotog».'),
    ('Asociación Porotog', 'ASOCIACIÓN POROTOG', None),
    ('Avellaneda', 'AVELLANEDA', 'Nombre oficial: Eliot Avellaneda.'),
]
PREDIOS_DISTINTOS = 0   # lo fija calcular()
CLASES = ['Pasto', 'Cebolla', 'Flores', 'Papa', 'Otros']
A_CLASE = {'Pasto no mejorado': 'Pasto', 'Pasto mejorado': 'Pasto', 'Cebolla': 'Cebolla',
           'Flores': 'Flores', 'Patrones de flores': 'Flores', 'Papas': 'Papa'}


def cargar():
    with io.open(os.path.join(GEO, 'fichas_predios.geojson'), encoding='utf-8') as f:
        fichas = [ft['properties'] for ft in json.load(f)['features']]
    with io.open(os.path.join(GEO, 'superficie_por_comunidad.json'), encoding='utf-8') as f:
        sup = {c['comunidad']: c for c in json.load(f)['comunidades']}
    with io.open(os.path.join(GEO, 'cultivos.json'), encoding='utf-8') as f:
        cultivos = json.load(f)
    with io.open(os.path.join(GEO, 'catastro_geo.geojson'), encoding='utf-8') as f:
        area_predio = {ft['properties']['clave_cata']: (ft['properties'].get('area_predi') or 0) / 10000.0
                       for ft in json.load(f)['features']}
    with io.open(os.path.join(GEO, 'codificacion_fichas.json'), encoding='utf-8') as f:
        codigo = json.load(f)['fichas']
    return fichas, sup, cultivos, area_predio, codigo


def clave_de(p):
    """Regla 14: clave_catastral manda, cod_poligono es el respaldo."""
    return (p.get('clave_catastral') or p.get('cod_poligono') or '').strip()


def sembrado_vs_predio(fichas, cultivos, area_predio, codigo):
    """Compara lo sembrado que se declaró en cada predio (sumando TODAS sus fichas,
    de cualquier comunidad) con lo que el predio mide en el catastro.

    El ajuste es proporcional por predio, igual que el riego ajustado: si lo
    declarado no cabe, cada cultivo de ese predio se reduce en la misma
    proporción hasta que la suma iguale la superficie catastral. Lo que ya cabe
    no se toca. Devuelve el factor por predio y el detalle de los que exceden."""
    ficha = {p['id']: p for p in fichas}
    sem = collections.defaultdict(float)
    fichas_de = collections.defaultdict(set)
    for c in cultivos:
        p = ficha.get(c['ficha_id'])
        if not p:
            continue
        k = clave_de(p)
        sem[k] += (c.get('superficie_m2') or 0) / 10000.0
        fichas_de[k].add(p['id'])
    factor, exceden = {}, []
    for k, s in sem.items():
        a = area_predio.get(k)
        if not a or s <= a:
            factor[k] = 1.0                       # sin polígono o cabe: no se ajusta
            continue
        factor[k] = a / s
        if s > 1.1 * a:                          # hasta 10 % de más se lee como redondeo
            fs = sorted(fichas_de[k], key=lambda i: codigo.get(i, ''))
            exceden.append(dict(clave=k, area=a, sembrado=s, veces=s / a,
                                comunidades=sorted({(ficha[i].get('comunidad') or '').strip() for i in fs}),
                                codigos=[codigo.get(i, '') for i in fs]))
    return factor, exceden


def calcular(fichas, sup, cultivos, area_predio, codigo):
    com_de = {p['id']: (p.get('comunidad') or '').strip() for p in fichas}
    clave_ficha = {p['id']: clave_de(p) for p in fichas}
    factor, exceden = sembrado_vs_predio(fichas, cultivos, area_predio, codigo)
    clase = collections.defaultdict(lambda: collections.defaultdict(float))   # com -> clase -> ha
    otros = collections.defaultdict(float)                                       # tipo -> ha (las 11)
    ajustada = collections.defaultdict(float)                                    # com -> ha ajustada
    nombres = {c for _, c, _ in COMUNIDADES}
    for c in cultivos:
        com = com_de.get(c['ficha_id'])
        if com not in nombres:
            continue
        ha = (c.get('superficie_m2') or 0) / 10000.0
        k = A_CLASE.get(c.get('tipo_cultivo'), 'Otros')
        clase[com][k] += ha
        ajustada[com] += ha * factor.get(clave_ficha.get(c['ficha_id']), 1.0)
        if k == 'Otros':
            otros[c.get('tipo_cultivo') or '(sin nombre del cultivo)'] += ha
    # predios con cultivo de cada comunidad y cuántos exceden (un predio cuenta en
    # cada comunidad donde tiene ficha)
    pred_com = collections.defaultdict(set)
    for c in cultivos:
        com = com_de.get(c['ficha_id'])
        if com in nombres:
            pred_com[com].add(clave_ficha.get(c['ficha_id']))
    exceden = [e for e in exceden if set(e['comunidades']) & nombres]
    global PREDIOS_DISTINTOS
    PREDIOS_DISTINTOS = len(set().union(*pred_com.values()))
    filas = []
    for etiqueta, com, nota in COMUNIDADES:
        assert com in sup, f'{com} no está en superficie_por_comunidad.json'
        fs = [p for p in fichas if (p.get('comunidad') or '').strip() == com]
        s = sup[com]
        filas.append(dict(etiqueta=etiqueta, com=com, nota=nota, sector=s.get('sector', ''),
                          principales=sum(1 for p in fs if not p.get('es_ficha_hija')),
                          adicionales=sum(1 for p in fs if p.get('es_ficha_hija')),
                          cat=s['superficie_catastral_ha'], riego=s['riego_ajustado_ha'],
                          sin_riego=s['sin_riego_catastral_ha'],
                          clases={k: clase[com].get(k, 0.0) for k in CLASES},
                          ajustada=ajustada[com], predios_cultivo=len(pred_com[com]),
                          exceden=sum(1 for e in exceden if com in e['comunidades']),
                          exceden_varias=sum(1 for e in exceden if com in e['comunidades']
                                             and len(e['codigos']) > 1)))
    exceden.sort(key=lambda e: (e['comunidades'][0], -e['veces']))
    return filas, dict(otros), exceden


# ── Excel ──────────────────────────────────────────────────────────────────
F = 'Arial'
AZUL = PatternFill('solid', fgColor='1F3864')
TOTAL = PatternFill('solid', fgColor='E2EFDA')
_s = Side(style='thin', color='BFBFBF')
BORDE = Border(left=_s, right=_s, top=_s, bottom=_s)
HA, ENT, PCT = '#,##0.00', '#,##0', '0.0%'


def _cab(ws, fila, textos, anchos):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t)
        c.font = Font(name=F, bold=True, color='FFFFFF', size=10); c.fill = AZUL; c.border = BORDE
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = anchos[j - 1]


def _c(ws, f, col, v, fmt=None, bold=False, fill=None):
    c = ws.cell(f, col, v); c.font = Font(name=F, size=10, bold=bold); c.border = BORDE
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c


def _titulo(ws, t, sub):
    ws['A1'] = t; ws['A1'].font = Font(name=F, bold=True, size=13, color='1F3864')
    ws['A2'] = sub; ws['A2'].font = Font(name=F, italic=True, size=9, color='595959')
    ws.sheet_view.showGridLines = False


def excel(filas, otros, exceden):
    wb = openpyxl.Workbook()
    corte = f'Datos al {FECHA_CORTE}. Fuente: padrón de usuarios de riego Guanguilquí–Porotog.'

    # 1. Superficie
    ws = wb.active; ws.title = 'Superficie'
    _titulo(ws, 'Superficie con riego y sin riego — 11 comunidades', corte + ' Superficie catastral: '
            'cada predio medido una vez por su polígono.')
    _cab(ws, 4, ['Comunidad', 'Sector', 'Fichas principales', 'Fichas adicionales', 'Total fichas',
                 'Superficie catastral (ha)', 'Con riego (ha)', 'Sin riego (ha)', '% con riego'],
         [22, 10, 12, 12, 10, 14, 12, 12, 10])
    r0 = 5
    for i, x in enumerate(filas, r0):
        _c(ws, i, 1, x['etiqueta'], bold=True); _c(ws, i, 2, x['sector'])
        _c(ws, i, 3, x['principales'], ENT); _c(ws, i, 4, x['adicionales'], ENT)
        _c(ws, i, 5, f'=C{i}+D{i}', ENT)
        _c(ws, i, 6, round(x['cat'], 2), HA); _c(ws, i, 7, round(x['riego'], 2), HA)
        _c(ws, i, 8, round(x['sin_riego'], 2), HA); _c(ws, i, 9, f'=IF(F{i}=0,0,G{i}/F{i})', PCT)
    rf = r0 + len(filas) - 1; t = rf + 1
    _c(ws, t, 1, 'Total 11 comunidades', bold=True, fill=TOTAL); _c(ws, t, 2, '', fill=TOTAL)
    for col, fmt in zip('CDEFGH', [ENT, ENT, ENT, HA, HA, HA]):
        _c(ws, t, openpyxl.utils.column_index_from_string(col), f'=SUM({col}{r0}:{col}{rf})', fmt, True, TOTAL)
    _c(ws, t, 9, f'=IF(F{t}=0,0,G{t}/F{t})', PCT, True, TOTAL)
    notas = [n_ for n_ in (f"{x['etiqueta']}: {x['nota']}" for x in filas if x['nota'])]
    for k, n_ in enumerate(notas + ['Con riego + sin riego = superficie catastral de cada comunidad.'], t + 2):
        ws.cell(k, 1, n_).font = Font(name=F, italic=True, size=9)
    ws.freeze_panes = 'B5'

    # 2. Área sembrada (las 5 clases)
    ws = wb.create_sheet('Área sembrada')
    _titulo(ws, 'Área sembrada por cultivo — 11 comunidades', corte + ' Superficie declarada en la sección '
            'de cultivos de las fichas.')
    _cab(ws, 4, ['Cultivo', 'Superficie actual (ha)', '% del total'], [22, 18, 12])
    for i, k in enumerate(CLASES, 5):
        _c(ws, i, 1, k, bold=True)
        _c(ws, i, 2, f"=INDEX('Cultivos por comunidad'!{openpyxl.utils.get_column_letter(2 + CLASES.index(k))}"
                     f"{5 + len(filas)},1)", HA)
        _c(ws, i, 3, f'=IF($B$10=0,0,B{i}/$B$10)', PCT)
    _c(ws, 10, 1, 'Total', bold=True, fill=TOTAL); _c(ws, 10, 2, '=SUM(B5:B9)', HA, True, TOTAL)
    _c(ws, 10, 3, '=SUM(C5:C9)', PCT, True, TOTAL)
    for k, n_ in enumerate([
            'Pasto = pasto mejorado + pasto no mejorado.',
            '«Otros» reúne el resto de cultivos y también bosque, monte y baldío declarados; ver la hoja «Detalle de Otros».',
            'Es superficie DECLARADA: no se compara con la catastral (en terrenos familiares varios herederos '
            'declaran el mismo terreno).'], 12):
        ws.cell(k, 1, n_).font = Font(name=F, italic=True, size=9)

    # 3. Cultivos por comunidad
    ws = wb.create_sheet('Cultivos por comunidad')
    _titulo(ws, 'Área sembrada por comunidad y cultivo (ha)', corte)
    _cab(ws, 4, ['Comunidad'] + CLASES + ['Total sembrado'], [22, 12, 12, 12, 12, 12, 14])
    for i, x in enumerate(filas, 5):
        _c(ws, i, 1, x['etiqueta'], bold=True)
        for j, k in enumerate(CLASES, 2):
            _c(ws, i, j, x['clases'][k], HA)
        _c(ws, i, 7, f'=SUM(B{i}:F{i})', HA, True)
    t = 5 + len(filas)
    _c(ws, t, 1, 'Total 11 comunidades', bold=True, fill=TOTAL)
    for j in range(2, 8):
        col = openpyxl.utils.get_column_letter(j)
        _c(ws, t, j, f'=SUM({col}5:{col}{t - 1})', HA, True, TOTAL)
    ws.freeze_panes = 'B5'

    # 4. Detalle de «Otros»
    ws = wb.create_sheet('Detalle de Otros')
    _titulo(ws, 'Qué hay dentro de «Otros» (11 comunidades)', corte)
    _cab(ws, 4, ['Cultivo o uso declarado', 'Superficie (ha)'], [30, 16])
    orden = sorted(otros.items(), key=lambda kv: -kv[1])
    for i, (k, v) in enumerate(orden, 5):
        _c(ws, i, 1, k); _c(ws, i, 2, v, HA)
    t = 5 + len(orden)
    _c(ws, t, 1, 'Total «Otros»', bold=True, fill=TOTAL); _c(ws, t, 2, f'=SUM(B5:B{t - 1})', HA, True, TOTAL)

    # 5. Sembrado frente al tamaño del predio
    ws = wb.create_sheet('Sembrado vs predio')
    _titulo(ws, 'Lo sembrado frente al tamaño del predio', corte + ' Ajustada: en los predios donde lo declarado '
            'no cabe, cada cultivo se reduce en la misma proporción hasta igualar la superficie catastral.')
    _cab(ws, 4, ['Comunidad', 'Sembrada declarada (ha)', 'Ajustada al predio (ha)', 'Diferencia (ha)',
                 '% de diferencia', 'Predios con cultivo', 'Predios que exceden más de 10 %',
                 'De ellos, con varias fichas'], [22, 14, 14, 12, 11, 12, 14, 13])
    for i, x in enumerate(filas, 5):
        _c(ws, i, 1, x['etiqueta'], bold=True); _c(ws, i, 2, sum(x['clases'].values()), HA)
        _c(ws, i, 3, x['ajustada'], HA); _c(ws, i, 4, f'=B{i}-C{i}', HA)
        _c(ws, i, 5, f'=IF(B{i}=0,0,D{i}/B{i})', PCT); _c(ws, i, 6, x['predios_cultivo'], ENT)
        _c(ws, i, 7, x['exceden'], ENT); _c(ws, i, 8, x['exceden_varias'], ENT)
    t = 5 + len(filas)
    _c(ws, t, 1, 'Total 11 comunidades', bold=True, fill=TOTAL)
    for col, fmt in zip('BCD', [HA, HA, HA]):
        _c(ws, t, openpyxl.utils.column_index_from_string(col), f'=SUM({col}5:{col}{t - 1})', fmt, True, TOTAL)
    _c(ws, t, 5, f'=IF(B{t}=0,0,D{t}/B{t})', PCT, True, TOTAL)
    _c(ws, t, 6, PREDIOS_DISTINTOS, ENT, True, TOTAL)   # distintos: no es la suma de la columna
    _c(ws, t, 7, len(exceden), ENT, True, TOTAL)
    _c(ws, t, 8, sum(1 for e in exceden if len(e['codigos']) > 1), ENT, True, TOTAL)
    for k, n_ in enumerate([
            'Un predio se compara con TODAS sus fichas: si varios herederos declaran el mismo terreno, sus cultivos '
            'se suman antes de compararlos con lo que mide el predio.',
            'Hasta un 10 % de más se considera redondeo y no se cuenta como exceso (sí se ajusta).',
            'Las columnas de predios no se suman entre comunidades: un predio con fichas de dos comunidades cuenta '
            'en ambas. Los totales de esas columnas son predios distintos.',
            'El detalle de cada predio que excede está en la hoja «Predios que exceden».'], t + 2):
        ws.cell(k, 1, n_).font = Font(name=F, italic=True, size=9)
    ws.freeze_panes = 'B5'

    # 6. Predios que exceden
    ws = wb.create_sheet('Predios que exceden')
    _titulo(ws, 'Predios cuyo sembrado declarado supera en más de 10 % lo que miden', corte + ' Los códigos '
            'de ficha permiten ubicar cada PDF en la carpeta de fichas.')
    _cab(ws, 4, ['Comunidad', 'Clave catastral', 'Área del predio (ha)', 'Sembrado declarado (ha)',
                 'Veces el predio', 'N.º de fichas', 'Códigos de ficha'], [22, 16, 12, 13, 10, 9, 60])
    for i, e in enumerate(exceden, 5):
        _c(ws, i, 1, ' / '.join(e['comunidades'])); c = _c(ws, i, 2, e['clave']); c.number_format = '@'
        _c(ws, i, 3, e['area'], HA); _c(ws, i, 4, e['sembrado'], HA)
        _c(ws, i, 5, f'=IF(C{i}=0,0,D{i}/C{i})', '0.0"×"'); _c(ws, i, 6, len(e['codigos']), ENT)
        c = _c(ws, i, 7, ', '.join(e['codigos'])); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.auto_filter.ref = f'A4:G{4 + len(exceden)}'; ws.freeze_panes = 'A5'

    wb.calculation.fullCalcOnLoad = True
    wb.save(XLSX)
    aplicar_formatos(XLSX)


# ── Word ───────────────────────────────────────────────────────────────────
AZUL_W = RGBColor(0x1F, 0x38, 0x64)


def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def _par(doc, texto, negrita=False, tam=10.5, color=None, cursiva=False, antes=0, despues=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(antes); p.paragraph_format.space_after = Pt(despues)
    r = p.add_run(texto); r.bold = negrita; r.italic = cursiva; r.font.size = Pt(tam)
    p.paragraph_format.keep_with_next = negrita or cursiva   # títulos y subtítulos van con su tabla
    if color: r.font.color.rgb = color
    return p


def _tabla(doc, cab, filas, anchos, total=None):
    t = doc.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(cab):
        c = t.rows[0].cells[j]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _sombrear(c, '1F3864')
    for k, fila in enumerate(filas + ([total] if total else [])):
        cells = t.add_row().cells
        es_total = total is not None and k == len(filas)
        for j, v in enumerate(fila):
            cells[j].text = ''
            r = cells[j].paragraphs[0].add_run(str(v)); r.font.size = Pt(8.5); r.bold = es_total or j == 0
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            if es_total: _sombrear(cells[j], 'E2EFDA')
    for fila in t.rows:
        for j, w in enumerate(anchos):
            fila.cells[j].width = Cm(w)
            for p in fila.cells[j].paragraphs: p.paragraph_format.space_after = Pt(0)
    # que la tabla no se parta entre páginas (son tablas cortas)
    for fila in t.rows[:-1]:
        for celda in fila.cells:
            for p in celda.paragraphs:
                p.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def word(filas, otros, exceden):
    doc = Document(MEMBRETE)
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    tot = lambda k: sum(x[k] for x in filas)  # noqa: E731
    sem = {k: sum(x['clases'][k] for x in filas) for k in CLASES}
    tot_sem = sum(sem.values())

    _par(doc, 'Superficie con riego, sin riego y área sembrada', True, 16, AZUL_W, despues=2)
    _par(doc, 'Once comunidades del sistema Porotog · datos al ' + FECHA_CORTE, tam=10, cursiva=True,
         color=RGBColor(0x59, 0x59, 0x59), despues=10)
    _par(doc, f"Este informe reúne {esn(tot('principales') + tot('adicionales'))} fichas "
              f"({esn(tot('principales'))} principales y {esn(tot('adicionales'))} adicionales) de once "
              f"comunidades: {', '.join(x['etiqueta'] for x in filas[:-1])} y {filas[-1]['etiqueta']}. "
              f"Presenta dos mediciones que conviene leer por separado: la superficie catastral, con su parte "
              f"con riego y sin riego, y el área sembrada que declararon los titulares en sus fichas.")

    _par(doc, '1. Superficie con riego y sin riego', True, 12, AZUL_W, antes=8)
    _par(doc, 'Superficie catastral: cada predio cuenta una sola vez, aunque varios herederos hayan declarado su '
              'parte por separado.', tam=9.5, cursiva=True)
    _tabla(doc, ['Comunidad', 'Fichas', 'Catastral (ha)', 'Con riego (ha)', 'Sin riego (ha)', '% con riego'],
           [[x['etiqueta'], esn(x['principales'] + x['adicionales']), esn(x['cat'], 2), esn(x['riego'], 2),
             esn(x['sin_riego'], 2), esn(100 * x['riego'] / x['cat'], 1) + ' %'] for x in filas],
           [4.2, 1.8, 2.6, 2.6, 2.6, 2.2],
           total=['Total', esn(tot('principales') + tot('adicionales')), esn(tot('cat'), 2), esn(tot('riego'), 2),
                  esn(tot('sin_riego'), 2), esn(100 * tot('riego') / tot('cat'), 1) + ' %'])

    _par(doc, '2. Área sembrada por cultivo', True, 12, AZUL_W, antes=8)
    _par(doc, 'Superficie declarada en la sección de cultivos de cada ficha, agrupada en cinco clases.',
         tam=9.5, cursiva=True)
    _tabla(doc, ['Cultivo', 'Superficie actual (ha)', '% del total'],
           [[k, esn(sem[k], 2), esn(100 * sem[k] / tot_sem, 1) + ' %'] for k in CLASES],
           [5, 4, 3], total=['Total', esn(tot_sem, 2), '100,0 %'])

    _par(doc, '3. Área sembrada por comunidad (ha)', True, 12, AZUL_W, antes=8)
    _tabla(doc, ['Comunidad'] + CLASES + ['Total'],
           [[x['etiqueta']] + [esn(x['clases'][k], 2) for k in CLASES] + [esn(sum(x['clases'].values()), 2)]
            for x in filas], [3.6, 2.1, 1.9, 1.7, 1.7, 1.9, 2.1],
           total=['Total'] + [esn(sem[k], 2) for k in CLASES] + [esn(tot_sem, 2)])

    tot_aj = sum(x['ajustada'] for x in filas)
    _par(doc, '4. Lo sembrado frente al tamaño del predio', True, 12, AZUL_W, antes=8)
    _par(doc, f'Muchas veces el área sembrada que declara el titular no cuadra con lo que mide su predio. De los '
              f'predios con cultivos de estas comunidades, {esn(len(exceden))} declaran más de lo que cabe (más de un '
              f'10 % por encima de su superficie catastral); en {esn(sum(1 for e in exceden if len(e["codigos"]) > 1))} '
              f'de ellos hay varias fichas sobre el mismo terreno. Si a cada predio se le limita lo sembrado a lo que '
              f'mide, las {esn(tot_sem, 2)} ha declaradas quedan en {esn(tot_aj, 2)} ha.', tam=10)
    _tabla(doc, ['Comunidad', 'Declarada (ha)', 'Ajustada al predio (ha)', 'Diferencia (ha)', 'Predios que exceden'],
           [[x['etiqueta'], esn(sum(x['clases'].values()), 2), esn(x['ajustada'], 2),
             esn(sum(x['clases'].values()) - x['ajustada'], 2), esn(x['exceden'])] for x in filas],
           [4.2, 2.8, 3.4, 2.8, 3.0],
           total=['Total', esn(tot_sem, 2), esn(tot_aj, 2), esn(tot_sem - tot_aj, 2), esn(len(exceden))])
    _par(doc, 'Por qué pasa:', True, 10, despues=2)
    for t_ in [
        'Herederos de un mismo terreno familiar: cada uno declara los cultivos de todo el terreno, y al sumarlos el '
        'predio parece sembrado dos o tres veces.',
        'Terreno arrendado o prestado fuera del predio: el titular siembra en otro lugar y lo declara junto con lo suyo.',
        'Área estimada a ojo en la entrevista, sin medir.',
    ]:
        p = _par(doc, '•  ' + t_, tam=9.5, despues=2)
        p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.35)
    _par(doc, 'Las tablas 2 y 3 muestran el área declarada, que es la que usa el análisis agrícola. La ajustada es '
              'una referencia de cuánto de eso cabe en los predios; el detalle de cada predio está en el Excel.',
         tam=9.5, cursiva=True, antes=4)

    _par(doc, 'Para leer bien estas cifras', True, 11, AZUL_W, antes=8)
    principales_otros = ', '.join(k.lower() for k, _ in sorted(otros.items(), key=lambda kv: -kv[1])[:6]
                                  if not k.startswith('('))
    for t_ in [
        'Pasto reúne el pasto mejorado y el no mejorado.',
        f'«Otros» reúne el resto de lo declarado, sobre todo {principales_otros}. El detalle completo está en el Excel.',
        'El área sembrada es declarada y la superficie catastral es medida: no se comparan en porcentaje. En los '
        'terrenos familiares varios herederos declaran el mismo terreno, por eso en algunas comunidades lo '
        'sembrado supera a la superficie catastral.',
    ] + [f"{x['etiqueta']}: {x['nota']}" for x in filas if x['nota']]:
        p = _par(doc, '•  ' + t_, tam=9.5, despues=3)
        p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.35)
    doc.save(DOCX)


def main():
    fichas, sup, cultivos, area_predio, codigo = cargar()
    filas, otros, exceden = calcular(fichas, sup, cultivos, area_predio, codigo)
    excel(filas, otros, exceden)
    word(filas, otros, exceden)
    print(f"Sembrada ajustada al predio {sum(x['ajustada'] for x in filas):.2f} ha · predios que exceden "
          f"más de 10 %: {len(exceden)} ({sum(1 for e in exceden if len(e['codigos']) > 1)} con varias fichas)")
    sem = sum(sum(x['clases'].values()) for x in filas)
    print(f"Fichas {sum(x['principales'] + x['adicionales'] for x in filas)} · catastral "
          f"{sum(x['cat'] for x in filas):.2f} ha · con riego {sum(x['riego'] for x in filas):.2f} · sin riego "
          f"{sum(x['sin_riego'] for x in filas):.2f} · sembrada {sem:.2f} ha")
    for k in CLASES:
        print(f"   {k:8} {sum(x['clases'][k] for x in filas):9.2f}")
    print('✔', XLSX); print('✔', DOCX)


if __name__ == '__main__':
    main()
