# -*- coding: utf-8 -*-
"""
Capítulo del informe técnico: "Los tres sectores de un vistazo".

Compara los tres sectores de investigación del sistema con barras agrupadas:
un gráfico por indicador, una barra por sector. Cierra el informe consolidado
(capítulo 7) y sirve suelto como entregable.

DE DÓNDE SALEN LAS CIFRAS
-------------------------
De `calcular_datos_por_corte()` del generador del informe por sector
(`generar_informe_sociologo_sector.py`), a través de `informe_graficos.
datos_dashboard()`. Es el MISMO cálculo que produce el informe por sector y
que reproduce el tablero web: aquí no se recalcula nada, solo se pone lado a
lado. Si una cifra de este capítulo no coincide con el informe 28, el error
está en uno de los dos scripts, no en los datos.

El sector de cada ficha es el de su comunidad según el catálogo oficial (el
mismo criterio de la web y, desde el 4-sep-2026, de todos los capítulos).

Encargo de JAVIKO, 4-sep-2026 (barras agrupadas, no los 11 gráficos por
sector repetidos: para eso está el informe 28).

SALIDAS
  docs/CAPITULO-los-tres-sectores.html
  build_entrega/Los_Tres_Sectores.xlsx
"""

import base64
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import informe_estilo as E  # noqa: E402
import informe_graficos as G  # noqa: E402
from generar_informe_sociologo import COLOR_SECTOR, RANGOS_PREDIO, f0, f1, f2, pct  # noqa: E402
from informe_estilo import esn

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
HTML = os.path.join(BASE, 'docs', 'CAPITULO-los-tres-sectores.html')
XLSX = os.path.join(BASE, 'build_entrega', 'Los_Tres_Sectores.xlsx')
MAPA = os.path.join(BASE, 'docs', 'graficos-sociologo', 'mapa-sistema.jpg')
SECTORES = ('Sector 1', 'Sector 2', 'Sector 3')


def resumen_por_sector(r):
    """Una fila por sector con los indicadores que se comparan."""
    datos, coms = r['datos'], r['comunidades']
    filas = {}
    for s in SECTORES:
        d = datos[s]
        cs = [c for c in coms if c['sector'] == s]
        n_pri = sum(c['n_pri'] for c in cs)
        con_viv = sum(c['con_vivienda'] for c in cs)
        ins = dict(d['instruccion'])
        ins_n = sum(ins.values()) or 1
        ten = dict(d['tenencia'])
        ten_n = sum(ten.values()) or 1
        com = {nom: (si, no) for nom, si, no in d['comunitaria']}
        filas[s] = {
            'comunidades': len(cs),
            'fichas': d['n_todas'], 'principales': d['n_pri'],
            'adicionales': d['n_todas'] - d['n_pri'],
            'predios': d['tamanos_total'],
            'declarada': sum(c['sup_declarada'] for c in cs),
            'declarada_riego': sum(c['sup_riego_decl'] for c in cs),
            'catastral': d['uso_suelo']['catastral'],
            'catastral_riego': d['uso_suelo']['riego'],
            'catastral_sin': d['uso_suelo']['sin_riego'],
            'caudal': sum(c['caudal_ls'] or 0 for c in cs if not c['caudal_heredado_de']),
            'instruccion_pct': {nv: pct(ins.get(nv, 0), ins_n) for nv in G.NIVELES_INSTRUCCION},
            'hijos_prom': d['hijos']['promedio'],
            'tenencia_pct': {k: pct(v, ten_n) for k, v in ten.items()},
            'metodo': {nom: v for nom, v, _ in d['metodo']},
            'comunitaria_pct': {k: pct(si, si + no) for k, (si, no) in com.items()},
            'destino': {nom: v for nom, v, _ in d['destino']},
            'destino_pct': {nom: pct(v, d['cultivos_registros']) for nom, v, _ in d['destino']},
            'cultivos': dict(d['cultivos_frec']),
            'pecuario': dict(d['pecuario']),
            'tamanos_pct': {et: pct(n, d['tamanos_total']) for et, n in d['tamanos']},
            'vivienda_pct': pct(con_viv, n_pri),
            'agua_pct': pct(sum(c['agua_consumo'] for c in cs), con_viv),
            'energia_pct': pct(sum(c['energia'] for c in cs), con_viv),
        }
    return filas


def serie(filas, f, fmt=None):
    """[(sector, color, [f(fila) por categoría])] en el orden de los sectores."""
    return [(s, COLOR_SECTOR[s], f(filas[s])) for s in SECTORES]


def main():
    G.preparar()
    G.configurar(os.path.join(BASE, 'docs'), 'graficos-capitulos')
    r = G.datos_dashboard()
    datos, sup, caudal = r['datos'], r['sup'], r['caudal']
    sis = datos['Todo el sistema']
    F = resumen_por_sector(r)
    tot = sup['total']

    # cuadre: los tres sectores deben sumar el sistema
    suma_f = sum(F[s]['fichas'] for s in SECTORES)
    assert suma_f == tot['fichas'], (suma_f, tot['fichas'])
    suma_p = sum(F[s]['predios'] for s in SECTORES)

    pct1 = lambda v: f'{esn(v, 1, False)} %'  # noqa: E731
    pct0 = lambda v: f'{esn(v, 0, False)} %'  # noqa: E731

    B = []
    A = B.append
    A(E.cabecera('Los tres sectores de un vistazo',
                 'Comparación de los sectores de investigación, indicador por '
                 'indicador · Capítulo del informe técnico'))
    A(E.aviso_corte(E.FECHA_CORTE, tot['regantes'], 0))
    A(E.kpis([
        (f0(F['Sector 1']['fichas']), f"fichas Sector 1 · {F['Sector 1']['comunidades']} comunidades"),
        (f0(F['Sector 2']['fichas']), f"fichas Sector 2 · {F['Sector 2']['comunidades']} comunidades"),
        (f0(F['Sector 3']['fichas']), f"fichas Sector 3 · {F['Sector 3']['comunidades']} comunidades"),
        (f2(caudal['totales']['caudal_sistema_ls']) + ' l/s', 'caudal del sistema'),
    ]))

    A('<h2>1. Cómo se organiza el sistema</h2>')
    A(f'<p>El sistema se investigó en <b>tres sectores</b> que agrupan a sus '
      f'{f0(len(r["comunidades"]))} comunidades. El Sector 1, con '
      f'{F["Sector 1"]["comunidades"]} comunidades y {f0(F["Sector 1"]["fichas"])} '
      'fichas, es el más numeroso; el Sector 3, con '
      f'{F["Sector 3"]["comunidades"]} comunidades y {f0(F["Sector 3"]["fichas"])} '
      f'fichas, es el de mayor superficie catastral por ficha. Cada ficha pertenece '
      'al sector de su comunidad según el listado oficial de organizaciones de '
      'riego, el mismo criterio de la aplicación web y de los demás capítulos.</p>')
    if os.path.exists(MAPA):
        with open(MAPA, 'rb') as f:
            b64 = base64.b64encode(f.read()).decode('ascii')
        A('<div class="fig evitar-corte"><img src="data:image/jpeg;base64,'
          f'{b64}" alt="Mapa de los sectores" style="width:100%"></div>')
        A('<p class="pie-fig">Comunas oficiales del GADM Cayambe recortadas al '
          'sistema y coloreadas por sector de investigación, sobre imagen '
          'satelital; la numeración 1–50 es la del listado oficial de '
          'organizaciones de riego.</p>')
    A('<table class="evitar-corte"><tr><th>Sector</th><th class="n">Comunidades</th>'
      '<th class="n">Fichas</th><th class="n">Principales</th>'
      '<th class="n">Adicionales</th><th class="n">Predios catastrales</th>'
      '<th class="n">Superficie catastral (ha)</th><th class="n">Caudal (l/s)</th></tr>')
    for s in SECTORES:
        x = F[s]
        A(f'<tr><td>{s}</td><td class="n">{x["comunidades"]}</td>'
          f'<td class="n">{f0(x["fichas"])}</td><td class="n">{f0(x["principales"])}</td>'
          f'<td class="n">{f0(x["adicionales"])}</td><td class="n">{f0(x["predios"])}</td>'
          f'<td class="n">{f2(x["catastral"])}</td><td class="n">{f1(x["caudal"])}</td></tr>')
    A(f'<tr class="dest"><td>Sistema</td><td class="n">{len(r["comunidades"])}</td>'
      f'<td class="n">{f0(tot["fichas"])}</td><td class="n">{f0(tot["regantes"])}</td>'
      f'<td class="n">{f0(tot["fichas"] - tot["regantes"])}</td>'
      f'<td class="n">{f0(tot["predios_catastrales"])}</td>'
      f'<td class="n">{f2(tot["superficie_catastral_ha"])}</td>'
      f'<td class="n">{f2(caudal["totales"]["caudal_comunidades_ls"])}</td></tr>')
    A('</table>')
    A('<p class="pie-fig">El caudal del sistema suma además '
      f'{f2(caudal["totales"]["caudal_individual_ls"])} l/s de '
      f'{caudal["totales"]["fichas_individuales"]} tomas individuales, que no '
      'se asignan a un sector.</p>')

    # ── padrón ──
    A('<h2>2. El padrón</h2>')
    b64 = G.g_barras_agrupadas(
        'sectores-padron', ['Fichas', 'Principales', 'Adicionales', 'Predios catastrales'],
        serie(F, lambda x: [x['fichas'], x['principales'], x['adicionales'], x['predios']]))
    A(E.figura(b64, 'Fichas y predios por sector',
               f'Todas las fichas del corte, {f0(tot["fichas"])}; predios catastrales '
               f'distintos, {f0(suma_p)}.'))

    # ── tierra ──
    A('<h2>3. La tierra</h2>')
    A('<p>Dos mediciones del mismo territorio, que no se suman entre sí: la '
      'catastral, que cuenta cada polígono municipal una sola vez, y la '
      'declarada por los titulares en la entrevista.</p>')
    b64 = G.g_barras_agrupadas(
        'sectores-catastral', ['Con riego', 'Sin riego', 'Total'],
        serie(F, lambda x: [round(x['catastral_riego']), round(x['catastral_sin']),
                            round(x['catastral'])]))
    A(E.figura(b64, 'Superficie catastral por sector (ha)',
               f'Medición catastral, {f2(tot["superficie_catastral_ha"])} ha en el '
               'sistema; riego ajustado al polígono.'))
    b64 = G.g_barras_agrupadas(
        'sectores-declarada', ['Con riego', 'Sin riego', 'Total'],
        serie(F, lambda x: [round(x['declarada_riego']),
                            round(x['declarada'] - x['declarada_riego']),
                            round(x['declarada'])]))
    A(E.figura(b64, 'Superficie declarada por sector (ha)',
               f'Declarada por los titulares, {f2(tot["superficie_declarada_ha"])} ha '
               'en el sistema.'))
    b64 = G.g_barras_agrupadas(
        'sectores-tamanos', [et for _, _, et in RANGOS_PREDIO],
        serie(F, lambda x: [round(x['tamanos_pct'][et], 1) for _, _, et in RANGOS_PREDIO]),
        fmt=pct0, rot=30)
    A(E.figura(b64, 'Tamaño de los predios por sector (% de los predios)',
               f'Predios catastrales, {f0(suma_p)}; cada uno contado una sola vez.'))

    # ── agua ──
    A('<h2>4. El agua</h2>')
    b64 = G.g_barras_v('sectores-caudal', [(s, F[s]['caudal']) for s in SECTORES],
                       [COLOR_SECTOR[s] for s in SECTORES], fmt=lambda v: f'{esn(v, 1)}')
    A(E.figura(b64, 'Caudal de las comunidades por sector (l/s)',
               'Un caudal por comunidad; '
               f'{f2(caudal["totales"]["caudal_comunidades_ls"])} l/s en el sistema.'))
    metodos = ['Aspersión', 'Gravedad', 'Goteo']
    b64 = G.g_barras_agrupadas(
        'sectores-metodo', metodos,
        serie(F, lambda x: [x['metodo'].get(m, 0) for m in metodos]), fmt=pct0)
    A(E.figura(b64, 'Método de riego por sector (promedio %)',
               'Todas las fichas; promedio simple del porcentaje declarado en cada '
               'ficha, redondeado a enteros.'))

    # ── personas ──
    A('<h2>5. Las personas</h2>')
    A('<p>Instrucción, familia, tenencia y conocimiento del proyecto se cuentan '
      'sobre las fichas principales: una por titular entrevistado.</p>')
    b64 = G.g_barras_agrupadas(
        'sectores-instruccion', G.NIVELES_INSTRUCCION,
        serie(F, lambda x: [round(x['instruccion_pct'][nv], 1) for nv in G.NIVELES_INSTRUCCION]),
        fmt=pct0)
    A(E.figura(b64, 'Nivel de instrucción por sector (% de los titulares)',
               f'Fichas principales con el dato, {f0(sis["instruccion_con_dato"])}.'))
    b64 = G.g_barras_v('sectores-hijos', [(s, round(F[s]['hijos_prom'], 2)) for s in SECTORES],
                       [COLOR_SECTOR[s] for s in SECTORES], fmt=lambda v: f'{esn(v, 1, False)}')
    A(E.figura(b64, 'Hijos por familia (promedio)',
               f'Fichas principales con el dato, {f0(sis["hijos"]["familias"])} familias.'))
    tenencias = [k for k, _ in sis['tenencia'] if k != 'Sin dato'][:4]
    b64 = G.g_barras_agrupadas(
        'sectores-tenencia', tenencias,
        serie(F, lambda x: [round(x['tenencia_pct'].get(t, 0), 1) for t in tenencias]),
        fmt=pct0, rot=15)
    A(E.figura(b64, 'Tenencia del predio por sector (% de los titulares)',
               f'Fichas principales, {f0(tot["regantes"])}.'))
    preguntas = [nom for nom, _, _ in sis['comunitaria']]
    b64 = G.g_barras_agrupadas(
        'sectores-comunitaria', preguntas,
        serie(F, lambda x: [round(x['comunitaria_pct'].get(q, 0), 1) for q in preguntas]),
        fmt=pct0)
    A(E.figura(b64, 'Represa y capacitación por sector (% que responde Sí)',
               'Fichas principales con respuesta.'))

    # ── producción ──
    A('<h2>6. La producción</h2>')
    top_c = [c for c, _ in sis['cultivos_frec'][:6]]
    b64 = G.g_barras_agrupadas(
        'sectores-cultivos', top_c,
        serie(F, lambda x: [x['cultivos'].get(c, 0) for c in top_c]), rot=20)
    A(E.figura(b64, 'Cultivos más frecuentes por sector (registros)',
               f'Todas las fichas; los seis cultivos con más registros del sistema, '
               f'de {f0(sis["cultivos_registros"])}.'))
    top_e = [e for e, _ in sis['pecuario'][:5]]
    b64 = G.g_barras_agrupadas(
        'sectores-pecuario', top_e,
        serie(F, lambda x: [x['pecuario'].get(e, 0) for e in top_e]), rot=15)
    A(E.figura(b64, 'Especies pecuarias principales por sector (cabezas)',
               'Todas las fichas; las cinco especies con más cabezas del sistema. '
               'No incluye la explotación avícola industrial de Asociación Rosalía.'))
    destinos = [nom for nom, _, _ in sis['destino']]
    b64 = G.g_barras_agrupadas(
        'sectores-destino', destinos,
        serie(F, lambda x: [round(x['destino_pct'].get(dd, 0), 1) for dd in destinos]),
        fmt=pct0)
    A(E.figura(b64, 'Destino de la producción agrícola por sector (% de los registros de cultivo)',
               'Todas las fichas; un cultivo puede declarar más de un destino.'))

    # ── vivienda ──
    A('<h2>7. La vivienda y sus servicios</h2>')
    b64 = G.g_barras_agrupadas(
        'sectores-vivienda', ['Con vivienda', 'Agua de consumo', 'Energía eléctrica'],
        serie(F, lambda x: [round(x['vivienda_pct'], 1), round(x['agua_pct'], 1),
                            round(x['energia_pct'], 1)]), fmt=pct0)
    A(E.figura(b64, 'Vivienda y servicios por sector (%)',
               'Con vivienda: % de las fichas principales. Agua y energía: % de las '
               'viviendas.'))

    # ── lectura ──
    A('<h2>8. Lectura</h2>')
    lider = max(SECTORES, key=lambda s: F[s]['catastral'])
    tec = {s: F[s]['metodo'].get('Aspersión', 0) + F[s]['metodo'].get('Goteo', 0)
           for s in SECTORES}
    presa = {s: F[s]['comunitaria_pct'].get('Conoce la represa', 0) for s in SECTORES}
    A('<ul>')
    A(f'<li>El <b>Sector 1</b> concentra {esn(pct(F["Sector 1"]["fichas"], tot["fichas"]), 0, False)} % '
      f'de las fichas y {esn(pct(F["Sector 1"]["caudal"], caudal["totales"]["caudal_comunidades_ls"]), 0, False)} % '
      'del caudal de las comunidades.</li>')
    A(f'<li>El <b>{lider}</b> es el de mayor superficie catastral '
      f'({f2(F[lider]["catastral"])} ha), con los predios más grandes.</li>')
    A(f'<li>La tecnificación del riego (aspersión más goteo, promedio por ficha) va de '
      f'{esn(min(tec.values()), 0, False)} % a {esn(max(tec.values()), 0, False)} % según el sector.</li>')
    A(f'<li>El conocimiento del proyecto de la represa es alto en los tres sectores: '
      f'entre {esn(min(presa.values()), 0, False)} % y {esn(max(presa.values()), 0, False)} %.</li>')
    A('</ul>')
    A(E.pie(E.FECHA_CORTE))

    os.makedirs(os.path.dirname(HTML), exist_ok=True)
    with open(HTML, 'w', encoding='utf-8') as f:
        f.write(E.documento('Los tres sectores de un vistazo — Padrón Guanguilquí–Porotog',
                            '\n'.join(B)))
    print(f'  capítulo: {os.path.relpath(HTML, BASE)}')

    # ── Excel ──
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    azul, blanco = PatternFill('solid', fgColor='1e4d8c'), Font(color='FFFFFF', bold=True)

    def hoja(nombre, cab, filas):
        ws = wb.create_sheet(nombre[:31])
        ws.append(cab)
        for c in ws[1]:
            c.fill, c.font = azul, blanco
        for f_ in filas:
            ws.append(f_)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                12, min(42, max(len(str(c.value or '')) for c in col) + 2))

    hoja('Resumen', ['Sector', 'Comunidades', 'Fichas', 'Principales', 'Adicionales',
                     'Predios catastrales', 'Catastral (ha)', 'Con riego catastral (ha)',
                     'Declarada (ha)', 'Con riego declarada (ha)', 'Caudal (l/s)'],
         [[s, F[s]['comunidades'], F[s]['fichas'], F[s]['principales'], F[s]['adicionales'],
           F[s]['predios'], round(F[s]['catastral'], 2), round(F[s]['catastral_riego'], 2),
           round(F[s]['declarada'], 2), round(F[s]['declarada_riego'], 2),
           round(F[s]['caudal'], 2)] for s in SECTORES])
    hoja('Instruccion %', ['Sector'] + G.NIVELES_INSTRUCCION,
         [[s] + [round(F[s]['instruccion_pct'][nv], 1) for nv in G.NIVELES_INSTRUCCION]
          for s in SECTORES])
    hoja('Tenencia %', ['Sector'] + tenencias,
         [[s] + [round(F[s]['tenencia_pct'].get(t, 0), 1) for t in tenencias] for s in SECTORES])
    hoja('Metodo %', ['Sector'] + metodos,
         [[s] + [F[s]['metodo'].get(m, 0) for m in metodos] for s in SECTORES])
    hoja('Represa y capacitacion %', ['Sector'] + preguntas,
         [[s] + [round(F[s]['comunitaria_pct'].get(q, 0), 1) for q in preguntas] for s in SECTORES])
    hoja('Cultivos (registros)', ['Sector'] + top_c,
         [[s] + [F[s]['cultivos'].get(c, 0) for c in top_c] for s in SECTORES])
    hoja('Pecuario (cabezas)', ['Sector'] + top_e,
         [[s] + [F[s]['pecuario'].get(e, 0) for e in top_e] for s in SECTORES])
    hoja('Destino %', ['Sector'] + destinos,
         [[s] + [round(F[s]['destino_pct'].get(dd, 0), 1) for dd in destinos] for s in SECTORES])
    hoja('Tamano predios %', ['Sector'] + [et for _, _, et in RANGOS_PREDIO],
         [[s] + [round(F[s]['tamanos_pct'][et], 1) for _, _, et in RANGOS_PREDIO]
          for s in SECTORES])
    hoja('Vivienda y servicios %', ['Sector', 'Con vivienda', 'Agua de consumo', 'Energia'],
         [[s, round(F[s]['vivienda_pct'], 1), round(F[s]['agua_pct'], 1),
           round(F[s]['energia_pct'], 1)] for s in SECTORES])
    hoja('Hijos', ['Sector', 'Promedio por familia'],
         [[s, round(F[s]['hijos_prom'], 2)] for s in SECTORES])
    del wb['Sheet']
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    from excel_compat import aplicar_formatos
    aplicar_formatos(XLSX)
    print(f'  excel   : {os.path.relpath(XLSX, BASE)}')
    for s in SECTORES:
        print(f'  {s}: {f0(F[s]["fichas"])} fichas · catastral {f2(F[s]["catastral"])} ha · '
              f'caudal {f1(F[s]["caudal"])} l/s · presa {esn(presa[s], 1, False)} %')


if __name__ == '__main__':
    main()
