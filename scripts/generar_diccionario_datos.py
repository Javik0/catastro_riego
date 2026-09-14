# -*- coding: utf-8 -*-
"""
Diccionario de datos del catastro socioeconómico y productivo predial.

Lo pide el Producto 5 del contrato: «adjuntar el diccionario de datos que
describa la estructura, campos, dominios y contenido de la información».

Describe las capas del GeoPackage del contratante (generar_gpkg_cliente.py),
que son las mismas que se entregan en Shapefile y en CAD. La estructura y los
conteos NO se escriben a mano: se leen del propio GeoPackage, y los dominios
de los campos categóricos se calculan consultando los valores que existen de
verdad. Así el diccionario no puede quedar desfasado del dato.

Salidas, dentro de la carpeta de entrega:
    0 - DICCIONARIO DE DATOS.pdf    documento formal
    DICCIONARIO DE DATOS.xlsx       la misma tabla, filtrable

Uso:
    python -X utf8 scripts/generar_diccionario_datos.py
"""
import os
import sqlite3
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from generar_fichas_pdf import (  # noqa: E402
    A4, AZUL, BORDE, TINTA, ANCHO_UTIL, ST_TD, ST_TH,
    cabecera_pie, colors, fmt_num, mm, Paragraph, ParagraphStyle, Spacer,
    Table, TableStyle, tabla_datos, titulo_seccion,
)
from informe_estilo import FECHA_CORTE  # noqa: E402
from reportlab.platypus import SimpleDocTemplate  # noqa: E402

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
GPKG = os.path.join(BASE, 'build_entrega', 'padron_riego_porotog.gpkg')
ENTREGA = r'C:\Users\HP\OneDrive\Escritorio\FICHAS PDF POROTOG'

# Orden de presentación y para qué sirve cada capa.
CAPAS = [
    ('predios_investigados', 'Polígono',
     'Predios con ficha de empadronamiento. Es la capa principal: un predio por '
     'fila, con el resumen de las fichas levantadas sobre él.'),
    ('fichas', 'Punto',
     'Una fila por ficha, ubicada en el punto GPS tomado en campo. Aquí está el '
     'detalle socioeconómico y productivo de cada titular.'),
    ('catastro_completo', 'Polígono',
     'Catastro rural del cantón, como contexto territorial. Incluye los predios '
     'sin ficha.'),
    ('comunas_oficiales', 'Polígono',
     'Límites de comunas según la capa oficial entregada por el contratante.'),
    ('sectores', 'Polígono', 'Los tres sectores de investigación.'),
    ('canales_riego', 'Línea', 'Red de conducción del sistema.'),
    ('cultivos', 'Tabla', 'Cultivos declarados en cada ficha (Sección 4 agrícola).'),
    ('animales', 'Tabla', 'Especies pecuarias declaradas en cada ficha (Sección 4).'),
]

# Campos cuyo dominio se calcula leyendo los valores existentes.
CATEGORICOS = {
    ('predios_investigados', 'estado_predio'), ('predios_investigados', 'condicion_riego'),
    ('fichas', 'tipo_ficha'), ('fichas', 'estado_investigacion'),
    ('fichas', 'tenencia_predio'), ('fichas', 'nivel_instruccion'),
    ('fichas', 'caudal_tipo'), ('fichas', 'frecuencia_riego'),
    ('fichas', 'tiene_reservorio'), ('fichas', 'tipo_tarifa'),
    ('fichas', 'material_construccion'), ('fichas', 'actividad_productiva'),
    ('catastro_completo', 'tiene_ficha'), ('comunas_oficiales', 'fuente'),
    ('sectores', 'sector'), ('cultivos', 'es_principal'),
    ('cultivos', 'destino'), ('animales', 'destino'),
}

SI_NO_VACIO = ('1 = sí · 0 = no · vacío = sin dato')

# Qué es cada campo. Sin esto el diccionario sería una lista de nombres.
DESCRIPCIONES = {
    'predios_investigados': {
        'clave_catastral': 'Clave catastral del predio en el catastro rural del GADM Cayambe. Identifica el polígono y enlaza con las demás capas.',
        'estado_predio': 'Si el predio se investigó con ficha principal o como predio adicional del titular.',
        'comunidad': 'Comunidad (organización de riego) a la que pertenece el predio.',
        'sector_riego': 'Sector de riego declarado en las fichas del predio.',
        'parroquia': 'Parroquia del cantón Cayambe donde está el predio.',
        'propietarios': 'Titulares con ficha en el predio. En el Shapefile este campo se llama TITULARES.',
        'propietario_catastro': 'Propietario según el catastro municipal. Puede no ser el titular entrevistado: el predio suele constar a nombre del padre o de «HEREDEROS DE…». En el Shapefile, PROP_CATAS.',
        'area_catastro_m2': 'Superficie del polígono según el catastro municipal, en metros cuadrados. Es la medición del territorio.',
        'total_fichas': 'Cuántas fichas se levantaron sobre este predio.',
        'fichas_principales': 'De ellas, cuántas son fichas principales (con entrevista).',
        'fichas_adicionales': 'De ellas, cuántas son fichas adicionales (otros predios del mismo titular).',
        'adicionales_pendientes': 'Fichas adicionales sin completar la Sección 4 (producción).',
        'area_declarada_m2': 'Superficie que declararon los titulares en las fichas del predio. No se mezcla con la catastral: son dos mediciones distintas.',
        'area_riego_m2': 'Superficie declarada bajo riego, en metros cuadrados.',
        'area_sin_riego_m2': 'Superficie declarada sin riego, en metros cuadrados.',
        'condicion_riego': 'Condición de riego del predio, calculada a partir de lo declarado.',
        'riego_pct': 'Porcentaje del área declarada que se riega. Sirve para leer los predios mixtos.',
        'caudal_comunidad_ls': 'Caudal en litros por segundo de la comunidad del predio. El caudal es por comunidad, no por ficha: no debe sumarse predio a predio.',
        'cultivos_predio': 'Cultivos declarados en el predio, resumidos en texto.',
        'animales_predio': 'Especies pecuarias declaradas en el predio, resumidas en texto.',
    },
    'fichas': {
        'ficha_id': 'Identificador interno de la ficha (UUID de QField). Enlaza con las tablas de cultivos y animales.',
        'codigo_ficha': 'CÓDIGO DEL PADRÓN: sector, comunidad, titular y ficha (S01-C22-R001-F01). Único por ficha, es el mismo con el que se nombra su ficha en PDF y el que muestra la plataforma web. En el Shapefile, codigo_fic.',
        'clave_catastral': 'Clave catastral del predio al que corresponde la ficha.',
        'codigo_predio': 'Código del formulario de campo (S-C-P001). Se repite entre fichas y no identifica a ninguna: se conserva solo por trazabilidad con QField.',
        'tipo_ficha': 'Si es la ficha principal del titular o una ficha adicional de otro predio suyo.',
        'estado_investigacion': 'Estado del levantamiento de la ficha.',
        'regante_principal': 'En una ficha adicional, el titular de su ficha principal.',
        'ficha_madre_id': 'En una ficha adicional, el identificador de su ficha principal.',
        'apellidos': 'Apellidos del titular entrevistado en campo.',
        'nombres': 'Nombres del titular entrevistado en campo.',
        'cedula': 'Cédula de identidad del titular. En organizaciones puede ser RUC.',
        'telefono_celular': 'Teléfono celular del titular. En el Shapefile, TEL_CEL.',
        'telefono_casa': 'Teléfono domiciliario del titular. En el Shapefile, TEL_CASA.',
        'parroquia': 'Parroquia donde está el predio de la ficha.',
        'comunidad': 'Comunidad (organización de riego) del titular.',
        'sector': 'Sector de riego declarado.',
        'sector_comunidad': 'Sector dentro de la comunidad, tal como lo nombran en campo.',
        'tenencia_predio': 'Forma de tenencia del predio declarada por el titular.',
        'nivel_instruccion': 'Nivel de instrucción del titular entrevistado.',
        'area_total_m2': 'Superficie total del predio declarada por el titular, en metros cuadrados.',
        'area_riego_m2': 'Superficie declarada bajo riego, en metros cuadrados.',
        'area_sin_riego_m2': 'Superficie declarada sin riego, en metros cuadrados.',
        'canal': 'Canal o ramal del que recibe el agua.',
        'caudal_ls': 'Caudal en litros por segundo. Cuando caudal_tipo dice «Recibe la Comunidad», es el caudal de la comunidad, no el de esta ficha.',
        'caudal_tipo': 'Si el caudal registrado es individual o el de la comunidad.',
        'frecuencia_riego': 'Cada cuánto le toca el turno de riego.',
        'dias_riego': 'Días de riego del turno.',
        'horas_turno': 'Horas que dura el turno.',
        'metodo_gravedad_pct': 'Porcentaje del riego por gravedad. En el Shapefile, metodo_gra.',
        'metodo_aspersion_pct': 'Porcentaje del riego por aspersión. En el Shapefile, metodo_asp.',
        'metodo_goteo_pct': 'Porcentaje del riego por goteo. En el Shapefile, metodo_got.',
        'valor_tarifa': 'Valor de la tarifa de riego, en dólares.',
        'tipo_tarifa': 'Periodicidad o modalidad de la tarifa.',
        'tiene_reservorio': 'Si el predio cuenta con reservorio y de qué tipo.',
        'agua_consumo': 'Si la vivienda dispone de agua de consumo. ' + SI_NO_VACIO,
        'energia_electrica': 'Si la vivienda dispone de energía eléctrica. ' + SI_NO_VACIO,
        'material_construccion': 'Material predominante de la vivienda.',
        'cota_msnm': 'Altitud del punto levantado, en metros sobre el nivel del mar.',
        'org_riego': 'Organización de riego a la que pertenece el titular.',
        'actividad_productiva': 'Actividad productiva principal declarada.',
        'observaciones': 'Observaciones registradas en campo. En Shapefile el texto se recorta a 254 caracteres por límite del formato; el texto completo consta en la ficha en PDF.',
        'investigado_por': 'Técnico investigador responsable del levantamiento. En las fichas adicionales generadas por el formulario, se hereda el técnico de la ficha principal.',
        'fecha_registro': 'Fecha en que se registró la ficha en campo.',
        'foto_url': 'Enlace a la fotografía tomada en campo, alojada en la plataforma.',
    },
    'catastro_completo': {
        'clave_catastral': 'Clave catastral del predio en el catastro rural del GADM Cayambe.',
        'propietario': 'Propietario según el catastro municipal, tal como lo entregó la entidad.',
        'cedula': 'Cédula del propietario según el catastro municipal.',
        'comunidad': 'Comunidad asignada al predio en el catastro municipal.',
        'area_predi_m2': 'Superficie del predio según el catastro municipal, en metros cuadrados.',
        'tiene_ficha': 'Si el predio tiene al menos una ficha de empadronamiento.',
    },
    'comunas_oficiales': {
        'comuna': 'Nombre de la comuna según la capa oficial. Se conservan las grafías de origen.',
        'area_comuna_ha': 'Superficie total de la comuna, en hectáreas.',
        'area_dentro_ha': 'Superficie de la comuna que cae dentro del área de estudio, en hectáreas.',
        'pct_dentro': 'Porcentaje de la comuna que cae dentro del área de estudio.',
        'fuente': 'Origen de la capa.',
    },
    'sectores': {
        'sector': 'Sector de investigación.',
        'total_fichas': 'Fichas levantadas en el sector.',
        'predios_catastro': 'Predios del catastro con ficha en el sector.',
        'area_dissolve_ha': 'Superficie del sector, en hectáreas, resultado de unir sus predios.',
        'area_riego_ha': 'Superficie bajo riego del sector, en hectáreas.',
        'caudal_total_ls': 'Caudal del sector, en litros por segundo.',
    },
    'canales_riego': {
        'nombre': 'Nombre de la red de conducción.',
    },
    'cultivos': {
        'ficha_id': 'Ficha en la que se declaró el cultivo. Enlaza con la capa fichas.',
        'clave_catastral': 'Clave catastral del predio de esa ficha.',
        'clave_predio_principal': 'Si la ficha es adicional, la clave del predio de su ficha principal.',
        'regante': 'Titular que declaró el cultivo.',
        'cultivo': 'Cultivo declarado.',
        'superficie_m2': 'Superficie sembrada declarada, en metros cuadrados.',
        'es_principal': 'Si es el cultivo principal del predio.',
        'destino': 'Destino de la producción. Un cultivo puede tener más de un destino.',
    },
    'animales': {
        'ficha_id': 'Ficha en la que se declaró la especie. Enlaza con la capa fichas.',
        'clave_catastral': 'Clave catastral del predio de esa ficha.',
        'clave_predio_principal': 'Si la ficha es adicional, la clave del predio de su ficha principal.',
        'regante': 'Titular que declaró la especie.',
        'especie': 'Especie pecuaria declarada.',
        'cantidad': 'Número de cabezas declaradas.',
        'destino': 'Destino de la producción. Una especie puede tener más de un destino.',
    },
}

TIPO_LEGIBLE = {'TEXT': 'Texto', 'REAL': 'Número decimal', 'INTEGER': 'Número entero'}


def leer_estructura():
    """Estructura y dominios, leídos del GeoPackage (no escritos a mano)."""
    con = sqlite3.connect(GPKG)
    cur = con.cursor()
    aux = con.cursor()   # cursor aparte: consultar con el mismo corta la iteración
    out = []
    for capa, geom, proposito in CAPAS:
        campos = []
        for r in cur.execute('PRAGMA table_info("{}")'.format(capa)).fetchall():
            nombre, tipo = r[1], r[2]
            if nombre in ('fid', 'geom'):
                continue
            dominio = ''
            if (capa, nombre) in CATEGORICOS:
                vals = [x[0] for x in aux.execute(
                    'SELECT DISTINCT "{}" FROM "{}" WHERE "{}" IS NOT NULL '
                    'AND TRIM("{}") <> \'\' ORDER BY 1'.format(nombre, capa, nombre, nombre))]
                if len(vals) <= 12:
                    dominio = ' · '.join(str(v) for v in vals)
                else:
                    dominio = '{} valores distintos'.format(fmt_num(len(vals)))
            elif nombre in ('agua_consumo', 'energia_electrica'):
                dominio = SI_NO_VACIO
            campos.append({
                'campo': nombre,
                'tipo': TIPO_LEGIBLE.get(tipo, tipo or 'Texto'),
                'descripcion': DESCRIPCIONES.get(capa, {}).get(nombre, ''),
                'dominio': dominio,
            })
        n = aux.execute('SELECT COUNT(*) FROM "{}"'.format(capa)).fetchone()[0]
        out.append({'capa': capa, 'geometria': geom, 'proposito': proposito,
                    'registros': n, 'campos': campos})
    con.close()
    return out


def escribir_excel(estructura, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = 'Diccionario'
    ws.append(['Capa', 'Geometría', 'Registros', 'Campo', 'Tipo',
               'Descripción', 'Dominio / valores admitidos'])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1E3A8A')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for capa in estructura:
        for i, f in enumerate(capa['campos']):
            ws.append([capa['capa'] if i == 0 else '',
                       capa['geometria'] if i == 0 else '',
                       capa['registros'] if i == 0 else '',
                       f['campo'], f['tipo'], f['descripcion'], f['dominio']])
    for col, ancho in zip('ABCDEFG', [24, 12, 11, 24, 16, 78, 44]):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)
    return ruta


def escribir_pdf(estructura, ruta):
    doc = SimpleDocTemplate(
        ruta, pagesize=A4, leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=27 * mm, bottomMargin=13 * mm,
        title='Diccionario de datos — Catastro socioeconómico y productivo predial',
        author='AP&CATASTROS — Padrón Guanguilquí–Porotog')

    st_h = ParagraphStyle('h', fontName='Helvetica-Bold', fontSize=12.5,
                          leading=15, textColor=AZUL, alignment=1)
    st_sub = ParagraphStyle('sub', fontName='Helvetica', fontSize=9.5, leading=12,
                            alignment=1, textColor=colors.HexColor('#475569'))
    st_p = ParagraphStyle('p', fontName='Helvetica', fontSize=9, leading=12.5,
                          textColor=TINTA, alignment=4)
    st_li = ParagraphStyle('li', parent=st_p, leftIndent=8 * mm, spaceAfter=2)

    h = [Paragraph('Diccionario de datos', st_h), Spacer(0, 1.5 * mm),
         Paragraph('Catastro socioeconómico y productivo predial · Sistema de riego '
                   'comunitario Guanguilquí–Porotog', st_sub),
         Spacer(0, 5 * mm),
         Paragraph('Describe la estructura, los campos, los dominios y el contenido de '
                   'la base de datos geoespacial del catastro. Las mismas capas se '
                   'entregan en GeoPackage (con el proyecto de QGIS ya simbolizado), en '
                   'Shapefile y en CAD. Los datos corresponden al levantamiento de campo '
                   f'cerrado al {FECHA_CORTE}.', st_p)]

    h += titulo_seccion('Las capas')
    filas = [[c['capa'], c['geometria'], fmt_num(c['registros']), c['proposito']]
             for c in estructura]
    h.append(tabla_datos(['Capa', 'Geometría', 'Registros', 'Qué contiene'], filas,
                         [40 * mm, 20 * mm, 22 * mm, 108 * mm]))

    h += titulo_seccion('Cómo se enlazan')
    h += [Paragraph('<b>clave_catastral</b> · enlaza predios_investigados, '
                    'catastro_completo y fichas. Es la clave del predio en el catastro '
                    'municipal.', st_li),
          Paragraph('<b>ficha_id</b> · enlaza cada ficha con sus cultivos y sus '
                    'animales.', st_li),
          Paragraph('<b>codigo_ficha</b> · identificador único de cada ficha '
                    '(S01-C22-R001-F01). Es el mismo con el que se nombra su ficha en '
                    'PDF, de modo que del mapa se llega al documento y al revés.', st_li),
          Spacer(0, 2 * mm),
          Paragraph('Un predio puede tener varias fichas —el terreno familiar declarado '
                    'por varios titulares—, así que entre predios_investigados y fichas '
                    'la relación es de uno a varios. El proyecto de QGIS que acompaña al '
                    'GeoPackage ya trae esa relación configurada: al hacer clic en un '
                    'predio se despliegan sus fichas, sus cultivos y sus animales.', st_p)]

    h += titulo_seccion('Dos mediciones de superficie que no se mezclan')
    h += [Paragraph('La <b>catastral</b> (area_catastro_m2) mide el polígono del catastro '
                    'municipal: cada predio cuenta una vez y es la superficie del '
                    'territorio. La <b>declarada</b> (area_declarada_m2, area_total_m2) '
                    'es la que informó cada titular en su ficha; en los predios de '
                    'herederos varios titulares declaran el mismo terreno familiar. Cada '
                    'cifra se lee dentro de su propia familia.', st_p)]

    for capa in estructura:
        h += titulo_seccion('Capa: {}'.format(capa['capa']))
        h += [Paragraph('{} · {} registros. {}'.format(
            capa['geometria'], fmt_num(capa['registros']), capa['proposito']), st_p),
            Spacer(0, 2 * mm)]
        filas = [[f['campo'], f['tipo'], f['descripcion'], f['dominio'] or '—']
                 for f in capa['campos']]
        h.append(tabla_datos(['Campo', 'Tipo', 'Descripción', 'Dominio'], filas,
                             [34 * mm, 22 * mm, 92 * mm, 42 * mm]))

    cab = {'creado_por': None, '_investigador': 'AP&CATASTROS'}
    doc.build(h, onFirstPage=lambda c, d: cabecera_pie(c, d, cab),
              onLaterPages=lambda c, d: cabecera_pie(c, d, cab))
    return ruta


def main():
    if not os.path.exists(GPKG):
        print('ERROR: no se encuentra {}'.format(GPKG))
        return 1
    estructura = leer_estructura()
    total = sum(len(c['campos']) for c in estructura)
    print('Leído del GeoPackage: {} capas, {} campos'.format(len(estructura), total))
    sin_desc = [(c['capa'], f['campo']) for c in estructura
                for f in c['campos'] if not f['descripcion']]
    if sin_desc:
        print('  ⚠ campos sin descripción: {}'.format(sin_desc))
    os.makedirs(ENTREGA, exist_ok=True)
    x = escribir_excel(estructura, os.path.join(ENTREGA, 'DICCIONARIO DE DATOS.xlsx'))
    p = escribir_pdf(estructura, os.path.join(ENTREGA, '0 - DICCIONARIO DE DATOS.pdf'))
    print('✔ {}'.format(x))
    print('✔ {}'.format(p))
    return 0


if __name__ == '__main__':
    sys.exit(main())
